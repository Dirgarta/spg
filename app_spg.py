"""
APLIKASI PENJUALAN SPG (online)
================================
Role:
  master : kelola toko, produk, akun SPG, lihat & filter semua penjualan, export
  spg    : input penjualan / sample di toko yang ditugaskan (tampilan simpel, HP-friendly)

Keamanan:
  - password di-hash dengan scrypt + salt acak (tidak disimpan apa adanya)
  - sesi disimpan sebagai hash di database, cookie HttpOnly + SameSite=Strict (+Secure di HTTPS)
  - proteksi CSRF (token per sesi), rate limit & penguncian akun saat brute force
  - 2FA (kode aplikasi Google Authenticator) opsional, wajib untuk master bila diaktifkan
  - CSP dengan nonce, HSTS, dan pemaksaan HTTPS saat dijalankan di internet
  - semua aktivitas tercatat di audit log (IP + perangkat)
"""
import base64
import hashlib
import hmac
import io
import json
import os
import re
import secrets
import sqlite3
import struct
import sys
import threading
import time
import unicodedata
import uuid
from datetime import date, datetime, timedelta, timezone
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse

APP_NAME = "Penjualan SPG"
APP_VERSION = "1.0"
BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = Path(os.environ.get("SPG_DATA_DIR", BASE_DIR / "data"))
DB_PATH = DATA_DIR / "spg.db"
BACKUP_DIR = DATA_DIR / "backup"

HOST = os.environ.get("SPG_HOST", "127.0.0.1")
PORT = int(os.environ.get("SPG_PORT", "8090"))
# dipercaya hanya bila aplikasi berada di belakang reverse proxy (Caddy/nginx)
TRUST_PROXY = os.environ.get("SPG_TRUST_PROXY", "0") == "1"
REQUIRE_HTTPS = os.environ.get("SPG_REQUIRE_HTTPS", "1") == "1"
DEV_MODE = os.environ.get("SPG_DEV", "0") == "1"          # matikan paksa-HTTPS saat uji lokal
TZ_OFFSET = int(os.environ.get("SPG_TZ_OFFSET", "7"))      # WIB = 7
SESSION_HOURS = {"master": 8, "spg": 12}
EDIT_WINDOW_MIN = int(os.environ.get("SPG_EDIT_MENIT", "60"))   # SPG boleh ubah/hapus entrinya sendiri
SATUAN = ("BOTOL", "DUS")
JENIS = ("JUAL", "SAMPLE")

TZ = timezone(timedelta(hours=TZ_OFFSET))


def now():
    return datetime.now(TZ)


def now_iso():
    return now().isoformat(timespec="seconds")


def today_str():
    return now().date().isoformat()


def new_uid():
    return uuid.uuid4().hex


# =====================================================================
# DATABASE
# =====================================================================
SCHEMA = """
PRAGMA journal_mode=WAL;
CREATE TABLE IF NOT EXISTS setting (key TEXT PRIMARY KEY, value TEXT);
CREATE TABLE IF NOT EXISTS pengguna (
    uid TEXT PRIMARY KEY,
    username TEXT NOT NULL UNIQUE COLLATE NOCASE,
    nama TEXT NOT NULL,
    role TEXT NOT NULL,
    pw_hash TEXT NOT NULL, pw_salt TEXT NOT NULL,
    harus_ganti INTEGER DEFAULT 1,
    totp TEXT DEFAULT '', totp_aktif INTEGER DEFAULT 0,
    aktif INTEGER DEFAULT 1,
    telepon TEXT DEFAULT '', catatan TEXT DEFAULT '',
    gagal INTEGER DEFAULT 0, kunci_sampai TEXT DEFAULT '',
    terakhir_login TEXT DEFAULT '', dibuat TEXT NOT NULL, diubah TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS toko (
    uid TEXT PRIMARY KEY, nama TEXT NOT NULL, wilayah TEXT DEFAULT '', alamat TEXT DEFAULT '',
    catatan TEXT DEFAULT '', aktif INTEGER DEFAULT 1, dibuat TEXT NOT NULL, diubah TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS produk (
    uid TEXT PRIMARY KEY, kode TEXT NOT NULL, nama TEXT NOT NULL,
    harga_botol REAL DEFAULT 0, harga_dus REAL DEFAULT 0, isi_dus INTEGER DEFAULT 12,
    aktif INTEGER DEFAULT 1, urut INTEGER DEFAULT 0, dibuat TEXT NOT NULL, diubah TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS spg_toko (
    spg_uid TEXT NOT NULL, toko_uid TEXT NOT NULL, dibuat TEXT NOT NULL,
    PRIMARY KEY (spg_uid, toko_uid));
CREATE TABLE IF NOT EXISTS penjualan (
    uid TEXT PRIMARY KEY,
    waktu TEXT NOT NULL,            -- waktu kirim (jam server, zona WIB)
    tanggal TEXT NOT NULL,
    spg_uid TEXT NOT NULL, toko_uid TEXT NOT NULL, produk_uid TEXT NOT NULL,
    jenis TEXT NOT NULL,            -- JUAL / SAMPLE
    satuan TEXT NOT NULL,           -- BOTOL / DUS
    qty REAL NOT NULL,
    isi_dus INTEGER DEFAULT 1,
    harga REAL DEFAULT 0,           -- harga per satuan yang dipakai saat transaksi
    total REAL DEFAULT 0,
    botol_setara REAL DEFAULT 0,
    catatan TEXT DEFAULT '',
    ip TEXT DEFAULT '', perangkat TEXT DEFAULT '',
    dihapus INTEGER DEFAULT 0, diubah TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS sesi (
    id TEXT PRIMARY KEY,            -- hash token
    pengguna_uid TEXT NOT NULL, csrf TEXT NOT NULL,
    dibuat TEXT NOT NULL, terakhir TEXT NOT NULL, kedaluwarsa TEXT NOT NULL,
    ip TEXT DEFAULT '', perangkat TEXT DEFAULT '');
CREATE TABLE IF NOT EXISTS log (
    uid TEXT PRIMARY KEY, waktu TEXT NOT NULL, pengguna_uid TEXT DEFAULT '',
    username TEXT DEFAULT '', role TEXT DEFAULT '', aksi TEXT NOT NULL,
    detail TEXT DEFAULT '', ip TEXT DEFAULT '', perangkat TEXT DEFAULT '');
CREATE INDEX IF NOT EXISTS ix_jual_tgl ON penjualan(tanggal);
CREATE INDEX IF NOT EXISTS ix_jual_spg ON penjualan(spg_uid);
CREATE INDEX IF NOT EXISTS ix_jual_toko ON penjualan(toko_uid);
CREATE INDEX IF NOT EXISTS ix_log_waktu ON log(waktu);
CREATE INDEX IF NOT EXISTS ix_sesi_user ON sesi(pengguna_uid);
"""

LOCK = threading.RLock()


def db():
    conn = sqlite3.connect(DB_PATH, timeout=20)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys=ON")
    return conn


def rows(cur):
    return [dict(r) for r in cur.fetchall()]


def one(cur):
    r = cur.fetchone()
    return dict(r) if r else None


class ApiError(Exception):
    def __init__(self, status, message, extra=None):
        super().__init__(message)
        self.status, self.message, self.extra = status, message, extra or {}


# =====================================================================
# PASSWORD & TOTP
# =====================================================================
SCRYPT = dict(n=2 ** 15, r=8, p=1, dklen=32, maxmem=64 * 1024 * 1024)
PW_UMUM = {"password", "12345678", "123456789", "qwerty123", "admin123", "spg12345",
           "password1", "rahasia123", "indonesia", "11223344", "abcd1234"}


def hash_pw(password, salt=None):
    salt = salt or secrets.token_hex(16)
    h = hashlib.scrypt(password.encode(), salt=bytes.fromhex(salt), **SCRYPT)
    return h.hex(), salt


def cek_pw(password, pw_hash, salt):
    try:
        h, _ = hash_pw(password, salt)
    except ValueError:
        return False
    return hmac.compare_digest(h, pw_hash)


def cek_kekuatan_pw(pw, username="", nama=""):
    if len(pw) < 10:
        raise ApiError(400, "Password minimal 10 karakter.")
    if len(pw) > 128:
        raise ApiError(400, "Password terlalu panjang.")
    if not re.search(r"[A-Za-z]", pw) or not re.search(r"\d", pw):
        raise ApiError(400, "Password harus berisi huruf dan angka.")
    low = pw.lower()
    if low in PW_UMUM:
        raise ApiError(400, "Password terlalu umum, gunakan yang lain.")
    if username and username.lower() in low:
        raise ApiError(400, "Password tidak boleh memuat username.")
    for bagian in (nama or "").lower().split():
        if len(bagian) >= 4 and bagian in low:
            raise ApiError(400, "Password tidak boleh memuat nama Anda.")
    if re.fullmatch(r"(.)\1+", pw) or low in "abcdefghijklmnopqrstuvwxyz" or low in "01234567890":
        raise ApiError(400, "Password terlalu mudah ditebak.")
    return True


def password_acak(n=12):
    huruf = "abcdefghijkmnopqrstuvwxyzABCDEFGHJKLMNPQRSTUVWXYZ"
    angka = "23456789"
    pool = huruf + angka
    while True:
        pw = "".join(secrets.choice(pool) for _ in range(n))
        if re.search(r"[A-Za-z]", pw) and re.search(r"\d", pw):
            return pw


def totp_secret():
    return base64.b32encode(secrets.token_bytes(20)).decode().rstrip("=")


def totp_cek(secret, kode, window=1):
    kode = re.sub(r"\D", "", kode or "")
    if len(kode) != 6 or not secret:
        return False
    key = base64.b32decode(secret + "=" * (-len(secret) % 8))
    counter = int(time.time()) // 30
    for delta in range(-window, window + 1):
        msg = struct.pack(">Q", counter + delta)
        h = hmac.new(key, msg, hashlib.sha1).digest()
        off = h[-1] & 0x0F
        val = (struct.unpack(">I", h[off:off + 4])[0] & 0x7FFFFFFF) % 1_000_000
        if hmac.compare_digest(f"{val:06d}", kode):
            return True
    return False


# =====================================================================
# SETTING, LOG, SESI
# =====================================================================
def get_set(conn, key, default=None):
    r = conn.execute("SELECT value FROM setting WHERE key=?", (key,)).fetchone()
    return r[0] if r else default


def set_set(conn, key, value):
    conn.execute("INSERT OR REPLACE INTO setting VALUES (?,?)", (key, str(value)))


def catat(conn, user, aksi, detail="", ip="", perangkat=""):
    conn.execute("INSERT INTO log VALUES (?,?,?,?,?,?,?,?,?)",
                 (new_uid(), now_iso(), (user or {}).get("uid", ""), (user or {}).get("username", ""),
                  (user or {}).get("role", ""), aksi, detail, ip, (perangkat or "")[:200]))


def buat_sesi(conn, user, ip, perangkat):
    token = secrets.token_urlsafe(32)
    sid = hashlib.sha256(token.encode()).hexdigest()
    csrf = secrets.token_urlsafe(24)
    jam = SESSION_HOURS.get(user["role"], 8)
    conn.execute("INSERT INTO sesi VALUES (?,?,?,?,?,?,?,?)",
                 (sid, user["uid"], csrf, now_iso(), now_iso(),
                  (now() + timedelta(hours=jam)).isoformat(timespec="seconds"), ip, (perangkat or "")[:200]))
    conn.execute("DELETE FROM sesi WHERE kedaluwarsa < ?", (now_iso(),))
    return token, csrf


def ambil_sesi(conn, token):
    if not token:
        return None
    sid = hashlib.sha256(token.encode()).hexdigest()
    s = one(conn.execute(
        """SELECT s.*, p.username, p.nama, p.role, p.aktif, p.harus_ganti, p.totp_aktif
           FROM sesi s JOIN pengguna p ON p.uid = s.pengguna_uid WHERE s.id=?""", (sid,)))
    if not s:
        return None
    if s["kedaluwarsa"] < now_iso() or not s["aktif"]:
        conn.execute("DELETE FROM sesi WHERE id=?", (sid,))
        return None
    # idle timeout 2 jam
    if datetime.fromisoformat(s["terakhir"]) < now() - timedelta(hours=2):
        conn.execute("DELETE FROM sesi WHERE id=?", (sid,))
        return None
    conn.execute("UPDATE sesi SET terakhir=? WHERE id=?", (now_iso(), sid))
    s["uid"] = s["pengguna_uid"]
    return s


# =====================================================================
# RATE LIMIT (per IP, dalam memori)
# =====================================================================
_hits = {}
_hits_lock = threading.Lock()


def rate_limit(ip, bucket, batas, detik):
    kunci = (ip, bucket)
    sekarang = time.time()
    with _hits_lock:
        arr = [t for t in _hits.get(kunci, []) if sekarang - t < detik]
        if len(arr) >= batas:
            arr.append(sekarang)
            _hits[kunci] = arr
            sisa = int(detik - (sekarang - arr[0]))
            raise ApiError(429, f"Terlalu banyak percobaan. Coba lagi dalam {max(sisa, 1)} detik.")
        arr.append(sekarang)
        _hits[kunci] = arr
        if len(_hits) > 5000:
            for k in [k for k, v in _hits.items() if not v or sekarang - v[-1] > 3600]:
                _hits.pop(k, None)


# =====================================================================
# HELPER
# =====================================================================
def bersih(v):
    return " ".join(str(v or "").split())


def upper(v):
    return bersih(v).upper()


def angka(v, nama="Angka"):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace("Rp", "")
    if re.fullmatch(r"-?\d{1,3}([.,]\d{3})+", s):
        s = re.sub(r"[.,]", "", s)
    elif "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        raise ApiError(400, f"{nama} tidak valid: {v}")


def tanggal(v, nama="Tanggal"):
    if not v:
        return None
    s = str(v).strip()
    for f in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, f).date().isoformat()
        except ValueError:
            pass
    raise ApiError(400, f"{nama} tidak valid: {v}")


def q1(q, k, d=""):
    return (q.get(k) or [d])[0].strip()


def slug_user(v):
    v = unicodedata.normalize("NFKD", bersih(v)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9._-]", "", v.replace(" ", "."))[:32]


def init_db():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    try:
        os.chmod(DATA_DIR, 0o700)
    except OSError:
        pass
    with db() as conn:
        conn.executescript(SCHEMA)
        if not conn.execute("SELECT 1 FROM pengguna WHERE role='master'").fetchone():
            pw = os.environ.get("SPG_ADMIN_PASSWORD") or password_acak(14)
            h, salt = hash_pw(pw)
            conn.execute("""INSERT INTO pengguna (uid, username, nama, role, pw_hash, pw_salt, harus_ganti,
                            dibuat, diubah) VALUES (?,?,?,?,?,?,1,?,?)""",
                         (new_uid(), "master", "Master", "master", h, salt, now_iso(), now_iso()))
            (DATA_DIR / "PASSWORD_AWAL.txt").write_text(
                f"Username : master\nPassword : {pw}\n\n"
                "Segera login dan ganti password ini. Hapus file ini setelah dicatat.\n", encoding="utf-8")
            print("=" * 64)
            print("  AKUN MASTER DIBUAT")
            print(f"  username : master")
            print(f"  password : {pw}")
            print(f"  (tersimpan juga di {DATA_DIR / 'PASSWORD_AWAL.txt'} — hapus setelah dicatat)")
            print("=" * 64)
        if not get_set(conn, "nama_perusahaan"):
            set_set(conn, "nama_perusahaan", os.environ.get("SPG_NAMA", "Penjualan SPG"))


def backup_harian():
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    target = BACKUP_DIR / f"spg_{today_str()}.db"
    if DB_PATH.exists() and not target.exists():
        src, dst = sqlite3.connect(DB_PATH), sqlite3.connect(target)
        src.backup(dst)
        dst.close(); src.close()
    for f in sorted(BACKUP_DIR.glob("spg_*.db"))[:-30]:
        f.unlink(missing_ok=True)


# =====================================================================
# LOGIN / AKUN
# =====================================================================
def login(body, ip, perangkat):
    rate_limit(ip, "login", 10, 300)
    username = bersih(body.get("username")).lower()
    password = str(body.get("password") or "")
    kode = bersih(body.get("kode"))
    if not username or not password:
        raise ApiError(400, "Username dan password wajib diisi.")
    with LOCK, db() as conn:
        u = one(conn.execute("SELECT * FROM pengguna WHERE username=? COLLATE NOCASE", (username,)))
        gagal_umum = ApiError(401, "Username atau password salah.")
        if not u:
            catat(conn, None, "LOGIN_GAGAL", f"username tidak ada: {username}", ip, perangkat)
            conn.commit()
            raise gagal_umum
        if not u["aktif"]:
            catat(conn, u, "LOGIN_GAGAL", "akun nonaktif", ip, perangkat)
            conn.commit()
            raise ApiError(403, "Akun ini dinonaktifkan. Hubungi master.")
        if u["kunci_sampai"] and u["kunci_sampai"] > now_iso():
            sisa = int((datetime.fromisoformat(u["kunci_sampai"]) - now()).total_seconds() / 60) + 1
            raise ApiError(429, f"Akun terkunci sementara karena salah password. Coba lagi {sisa} menit lagi.")
        if not cek_pw(password, u["pw_hash"], u["pw_salt"]):
            gagal = u["gagal"] + 1
            kunci = (now() + timedelta(minutes=15)).isoformat(timespec="seconds") if gagal >= 5 else ""
            conn.execute("UPDATE pengguna SET gagal=?, kunci_sampai=? WHERE uid=?",
                         (0 if kunci else gagal, kunci, u["uid"]))
            catat(conn, u, "LOGIN_GAGAL", f"password salah (ke-{gagal})" + (" — akun dikunci 15 menit" if kunci else ""), ip, perangkat)
            conn.commit()
            if kunci:
                raise ApiError(429, "Password salah 5 kali. Akun dikunci 15 menit.")
            raise gagal_umum
        if u["totp_aktif"]:
            if not kode:
                raise ApiError(401, "Masukkan kode 6 angka dari aplikasi Authenticator.", {"butuh_kode": True})
            if not totp_cek(u["totp"], kode):
                catat(conn, u, "LOGIN_GAGAL", "kode 2FA salah", ip, perangkat)
                conn.commit()
                raise ApiError(401, "Kode 2FA salah.", {"butuh_kode": True})
        token, csrf = buat_sesi(conn, u, ip, perangkat)
        conn.execute("UPDATE pengguna SET gagal=0, kunci_sampai='', terakhir_login=? WHERE uid=?", (now_iso(), u["uid"]))
        catat(conn, u, "LOGIN", "berhasil masuk", ip, perangkat)
    return token, {"uid": u["uid"], "username": u["username"], "nama": u["nama"], "role": u["role"],
                   "harus_ganti": bool(u["harus_ganti"]), "csrf": csrf, "totp_aktif": bool(u["totp_aktif"])}


def logout(user, token, ip):
    with LOCK, db() as conn:
        if token:
            conn.execute("DELETE FROM sesi WHERE id=?", (hashlib.sha256(token.encode()).hexdigest(),))
        catat(conn, user, "LOGOUT", "", ip)
    return {"ok": True}


def ganti_password(user, body, ip):
    lama, baru = str(body.get("lama") or ""), str(body.get("baru") or "")
    with LOCK, db() as conn:
        u = one(conn.execute("SELECT * FROM pengguna WHERE uid=?", (user["uid"],)))
        if not cek_pw(lama, u["pw_hash"], u["pw_salt"]):
            rate_limit(ip, "gantipw", 10, 600)
            raise ApiError(400, "Password lama salah.")
        if lama == baru:
            raise ApiError(400, "Password baru harus berbeda dengan yang lama.")
        cek_kekuatan_pw(baru, u["username"], u["nama"])
        h, salt = hash_pw(baru)
        conn.execute("UPDATE pengguna SET pw_hash=?, pw_salt=?, harus_ganti=0, diubah=? WHERE uid=?",
                     (h, salt, now_iso(), u["uid"]))
        # keluarkan semua sesi lain
        conn.execute("DELETE FROM sesi WHERE pengguna_uid=? AND id!=?",
                     (u["uid"], hashlib.sha256((body.get("_token") or "").encode()).hexdigest()))
        catat(conn, user, "GANTI_PASSWORD", "password diganti sendiri", ip)
    return {"ok": True}


def totp_siapkan(user):
    secret = totp_secret()
    with LOCK, db() as conn:
        conn.execute("UPDATE pengguna SET totp=?, totp_aktif=0 WHERE uid=?", (secret, user["uid"]))
        nama = get_set(conn, "nama_perusahaan", APP_NAME)
    uri = f"otpauth://totp/{nama}:{user['username']}?secret={secret}&issuer={nama}&digits=6&period=30"
    return {"secret": secret, "uri": uri}


def totp_aktifkan(user, body, ip):
    with LOCK, db() as conn:
        u = one(conn.execute("SELECT totp FROM pengguna WHERE uid=?", (user["uid"],)))
        if not u["totp"] or not totp_cek(u["totp"], body.get("kode")):
            raise ApiError(400, "Kode salah. Pastikan jam HP sudah otomatis/tepat.")
        conn.execute("UPDATE pengguna SET totp_aktif=1, diubah=? WHERE uid=?", (now_iso(), user["uid"]))
        catat(conn, user, "2FA_AKTIF", "verifikasi 2 langkah diaktifkan", ip)
    return {"ok": True}


def totp_matikan(user, body, ip):
    with LOCK, db() as conn:
        u = one(conn.execute("SELECT * FROM pengguna WHERE uid=?", (user["uid"],)))
        if not cek_pw(str(body.get("password") or ""), u["pw_hash"], u["pw_salt"]):
            rate_limit(ip, "2fa", 10, 600)
            raise ApiError(400, "Password salah.")
        conn.execute("UPDATE pengguna SET totp='', totp_aktif=0, diubah=? WHERE uid=?", (now_iso(), user["uid"]))
        catat(conn, user, "2FA_MATI", "verifikasi 2 langkah dimatikan", ip)
    return {"ok": True}


def sesi_saya(user):
    with db() as conn:
        return rows(conn.execute(
            "SELECT dibuat, terakhir, kedaluwarsa, ip, perangkat FROM sesi WHERE pengguna_uid=? ORDER BY terakhir DESC",
            (user["uid"],)))


# =====================================================================
# MASTER: PENGGUNA (SPG)
# =====================================================================
def user_list():
    with db() as conn:
        data = rows(conn.execute(
            """SELECT p.uid, p.username, p.nama, p.role, p.aktif, p.telepon, p.catatan, p.harus_ganti,
                      p.totp_aktif, p.terakhir_login, p.kunci_sampai, p.dibuat,
                      (SELECT COUNT(*) FROM spg_toko st WHERE st.spg_uid=p.uid) AS jumlah_toko,
                      (SELECT COUNT(*) FROM penjualan j WHERE j.spg_uid=p.uid AND j.dihapus=0) AS jumlah_entri
               FROM pengguna p ORDER BY p.role DESC, p.nama"""))
        toko = {}
        for r in conn.execute("""SELECT st.spg_uid, t.uid, t.nama FROM spg_toko st JOIN toko t ON t.uid=st.toko_uid
                                 WHERE t.aktif=1 ORDER BY t.nama"""):
            toko.setdefault(r["spg_uid"], []).append({"uid": r["uid"], "nama": r["nama"]})
    for u in data:
        u["toko"] = toko.get(u["uid"], [])
    return data


def user_simpan(user, body, uid=None, ip=""):
    nama = bersih(body.get("nama"))
    role = body.get("role") if body.get("role") in ("master", "spg") else "spg"
    if not nama:
        raise ApiError(400, "Nama wajib diisi.")
    username = slug_user(body.get("username") or nama)
    if len(username) < 3:
        raise ApiError(400, "Username minimal 3 karakter (huruf/angka).")
    hasil = {}
    with LOCK, db() as conn:
        bentrok = conn.execute("SELECT 1 FROM pengguna WHERE username=? COLLATE NOCASE AND uid!=?",
                               (username, uid or "")).fetchone()
        if bentrok:
            raise ApiError(409, f"Username '{username}' sudah dipakai.")
        if uid:
            u = one(conn.execute("SELECT * FROM pengguna WHERE uid=?", (uid,)))
            if not u:
                raise ApiError(404, "Pengguna tidak ditemukan.")
            aktif = 1 if body.get("aktif", True) else 0
            if u["role"] == "master" and (role != "master" or not aktif):
                sisa = conn.execute("SELECT COUNT(*) FROM pengguna WHERE role='master' AND aktif=1 AND uid!=?", (uid,)).fetchone()[0]
                if not sisa:
                    raise ApiError(400, "Minimal harus ada satu master aktif.")
            conn.execute("""UPDATE pengguna SET username=?, nama=?, role=?, aktif=?, telepon=?, catatan=?, diubah=?
                            WHERE uid=?""",
                         (username, nama, role, aktif, bersih(body.get("telepon")), bersih(body.get("catatan")),
                          now_iso(), uid))
            if not aktif:
                conn.execute("DELETE FROM sesi WHERE pengguna_uid=?", (uid,))
            catat(conn, user, "UBAH_PENGGUNA", f"{username} ({role}){'' if aktif else ' — dinonaktifkan'}", ip)
        else:
            uid = new_uid()
            pw = password_acak()
            h, salt = hash_pw(pw)
            conn.execute("""INSERT INTO pengguna (uid, username, nama, role, pw_hash, pw_salt, harus_ganti, aktif,
                            telepon, catatan, dibuat, diubah) VALUES (?,?,?,?,?,?,1,1,?,?,?,?)""",
                         (uid, username, nama, role, h, salt, bersih(body.get("telepon")),
                          bersih(body.get("catatan")), now_iso(), now_iso()))
            hasil["password"] = pw
            catat(conn, user, "TAMBAH_PENGGUNA", f"{username} ({role})", ip)
        if "toko_uids" in body:
            _set_toko(conn, uid, body.get("toko_uids") or [])
    hasil["uid"] = uid
    hasil["username"] = username
    return hasil


def _set_toko(conn, spg_uid, uids):
    conn.execute("DELETE FROM spg_toko WHERE spg_uid=?", (spg_uid,))
    for t in uids:
        if conn.execute("SELECT 1 FROM toko WHERE uid=?", (t,)).fetchone():
            conn.execute("INSERT OR IGNORE INTO spg_toko VALUES (?,?,?)", (spg_uid, t, now_iso()))


def user_toko(user, uid, body, ip):
    with LOCK, db() as conn:
        if not conn.execute("SELECT 1 FROM pengguna WHERE uid=?", (uid,)).fetchone():
            raise ApiError(404, "Pengguna tidak ditemukan.")
        _set_toko(conn, uid, body.get("toko_uids") or [])
        n = conn.execute("SELECT COUNT(*) FROM spg_toko WHERE spg_uid=?", (uid,)).fetchone()[0]
        nm = one(conn.execute("SELECT username FROM pengguna WHERE uid=?", (uid,)))["username"]
        catat(conn, user, "SET_TOKO_SPG", f"{nm}: {n} toko", ip)
    return {"ok": True, "jumlah": n}


def user_reset_pw(user, uid, ip):
    pw = password_acak()
    h, salt = hash_pw(pw)
    with LOCK, db() as conn:
        u = one(conn.execute("SELECT username FROM pengguna WHERE uid=?", (uid,)))
        if not u:
            raise ApiError(404, "Pengguna tidak ditemukan.")
        conn.execute("UPDATE pengguna SET pw_hash=?, pw_salt=?, harus_ganti=1, gagal=0, kunci_sampai='', diubah=? WHERE uid=?",
                     (h, salt, now_iso(), uid))
        conn.execute("DELETE FROM sesi WHERE pengguna_uid=?", (uid,))
        catat(conn, user, "RESET_PASSWORD", f"password {u['username']} direset", ip)
    return {"password": pw, "username": u["username"]}


def user_buka_kunci(user, uid, ip):
    with LOCK, db() as conn:
        conn.execute("UPDATE pengguna SET gagal=0, kunci_sampai='' WHERE uid=?", (uid,))
        catat(conn, user, "BUKA_KUNCI", uid, ip)
    return {"ok": True}


def user_hapus(user, uid, ip):
    with LOCK, db() as conn:
        u = one(conn.execute("SELECT * FROM pengguna WHERE uid=?", (uid,)))
        if not u:
            raise ApiError(404, "Pengguna tidak ditemukan.")
        if u["role"] == "master" and conn.execute(
                "SELECT COUNT(*) FROM pengguna WHERE role='master' AND aktif=1 AND uid!=?", (uid,)).fetchone()[0] == 0:
            raise ApiError(400, "Minimal harus ada satu master aktif.")
        n = conn.execute("SELECT COUNT(*) FROM penjualan WHERE spg_uid=? AND dihapus=0", (uid,)).fetchone()[0]
        if n:
            conn.execute("UPDATE pengguna SET aktif=0, diubah=? WHERE uid=?", (now_iso(), uid))
            conn.execute("DELETE FROM sesi WHERE pengguna_uid=?", (uid,))
            catat(conn, user, "NONAKTIF_PENGGUNA", f"{u['username']} punya {n} entri, akun dinonaktifkan", ip)
            return {"ok": True, "dinonaktifkan": True, "entri": n}
        conn.execute("DELETE FROM spg_toko WHERE spg_uid=?", (uid,))
        conn.execute("DELETE FROM sesi WHERE pengguna_uid=?", (uid,))
        conn.execute("DELETE FROM pengguna WHERE uid=?", (uid,))
        catat(conn, user, "HAPUS_PENGGUNA", u["username"], ip)
    return {"ok": True}


# =====================================================================
# MASTER: TOKO & PRODUK
# =====================================================================
def toko_list(q=None):
    q = q or {}
    where, args = [], []
    if q1(q, "aktif", "1") == "1":
        where.append("t.aktif=1")
    if q1(q, "q"):
        where.append("(t.nama LIKE ? OR t.wilayah LIKE ? OR t.alamat LIKE ?)")
        args += [f"%{q1(q, 'q')}%"] * 3
    w = (" WHERE " + " AND ".join(where)) if where else ""
    with db() as conn:
        return rows(conn.execute(
            f"""SELECT t.*, (SELECT COUNT(*) FROM spg_toko st WHERE st.toko_uid=t.uid) AS jumlah_spg,
                       (SELECT COUNT(*) FROM penjualan j WHERE j.toko_uid=t.uid AND j.dihapus=0) AS jumlah_entri,
                       (SELECT group_concat(p.nama, ', ') FROM spg_toko st JOIN pengguna p ON p.uid=st.spg_uid
                        WHERE st.toko_uid=t.uid) AS spg
                FROM toko t {w} ORDER BY t.nama""", args))


def toko_simpan(user, body, uid=None, ip=""):
    nama = upper(body.get("nama"))
    if not nama:
        raise ApiError(400, "Nama toko wajib diisi.")
    data = (nama, upper(body.get("wilayah")), bersih(body.get("alamat")), bersih(body.get("catatan")),
            1 if body.get("aktif", True) else 0, now_iso())
    with LOCK, db() as conn:
        if conn.execute("SELECT 1 FROM toko WHERE nama=? COLLATE NOCASE AND uid!=?", (nama, uid or "")).fetchone():
            raise ApiError(409, f"Toko '{nama}' sudah ada.")
        if uid:
            conn.execute("UPDATE toko SET nama=?, wilayah=?, alamat=?, catatan=?, aktif=?, diubah=? WHERE uid=?",
                         (*data, uid))
            catat(conn, user, "UBAH_TOKO", nama, ip)
        else:
            uid = new_uid()
            conn.execute("INSERT INTO toko (uid, nama, wilayah, alamat, catatan, aktif, diubah, dibuat) VALUES (?,?,?,?,?,?,?,?)",
                         (uid, *data, now_iso()))
            catat(conn, user, "TAMBAH_TOKO", nama, ip)
        if body.get("spg_uids") is not None:
            conn.execute("DELETE FROM spg_toko WHERE toko_uid=?", (uid,))
            for s in body["spg_uids"]:
                conn.execute("INSERT OR IGNORE INTO spg_toko VALUES (?,?,?)", (s, uid, now_iso()))
    return {"uid": uid, "nama": nama}


def toko_hapus(user, uid, ip):
    with LOCK, db() as conn:
        t = one(conn.execute("SELECT * FROM toko WHERE uid=?", (uid,)))
        if not t:
            raise ApiError(404, "Toko tidak ditemukan.")
        n = conn.execute("SELECT COUNT(*) FROM penjualan WHERE toko_uid=? AND dihapus=0", (uid,)).fetchone()[0]
        if n:
            conn.execute("UPDATE toko SET aktif=0, diubah=? WHERE uid=?", (now_iso(), uid))
            catat(conn, user, "NONAKTIF_TOKO", f"{t['nama']} ({n} entri)", ip)
            return {"ok": True, "dinonaktifkan": True, "entri": n}
        conn.execute("DELETE FROM spg_toko WHERE toko_uid=?", (uid,))
        conn.execute("DELETE FROM toko WHERE uid=?", (uid,))
        catat(conn, user, "HAPUS_TOKO", t["nama"], ip)
    return {"ok": True}


def produk_list(aktif_saja=True):
    with db() as conn:
        return rows(conn.execute(
            f"""SELECT p.*, (SELECT COUNT(*) FROM penjualan j WHERE j.produk_uid=p.uid AND j.dihapus=0) AS jumlah_entri
                FROM produk p {'WHERE p.aktif=1' if aktif_saja else ''} ORDER BY p.urut, p.nama"""))


def produk_simpan(user, body, uid=None, ip=""):
    kode, nama = upper(body.get("kode")), bersih(body.get("nama"))
    if not kode or not nama:
        raise ApiError(400, "Kode dan nama produk wajib diisi.")
    hb, hd = angka(body.get("harga_botol"), "Harga botol"), angka(body.get("harga_dus"), "Harga dus")
    isi = int(angka(body.get("isi_dus"), "Isi per dus") or 12)
    if isi < 1:
        raise ApiError(400, "Isi per dus minimal 1.")
    if not hd and hb:
        hd = hb * isi
    with LOCK, db() as conn:
        if conn.execute("SELECT 1 FROM produk WHERE kode=? COLLATE NOCASE AND uid!=?", (kode, uid or "")).fetchone():
            raise ApiError(409, f"Kode produk '{kode}' sudah ada.")
        vals = (kode, nama, hb, hd, isi, 1 if body.get("aktif", True) else 0,
                int(angka(body.get("urut"))), now_iso())
        if uid:
            conn.execute("""UPDATE produk SET kode=?, nama=?, harga_botol=?, harga_dus=?, isi_dus=?, aktif=?, urut=?,
                            diubah=? WHERE uid=?""", (*vals, uid))
            catat(conn, user, "UBAH_PRODUK", f"{kode} {nama} — botol {hb:,.0f} / dus {hd:,.0f}", ip)
        else:
            uid = new_uid()
            conn.execute("""INSERT INTO produk (uid, kode, nama, harga_botol, harga_dus, isi_dus, aktif, urut, diubah, dibuat)
                            VALUES (?,?,?,?,?,?,?,?,?,?)""", (uid, *vals, now_iso()))
            catat(conn, user, "TAMBAH_PRODUK", f"{kode} {nama}", ip)
    return {"uid": uid}


def produk_hapus(user, uid, ip):
    with LOCK, db() as conn:
        p = one(conn.execute("SELECT * FROM produk WHERE uid=?", (uid,)))
        if not p:
            raise ApiError(404, "Produk tidak ditemukan.")
        n = conn.execute("SELECT COUNT(*) FROM penjualan WHERE produk_uid=? AND dihapus=0", (uid,)).fetchone()[0]
        if n:
            conn.execute("UPDATE produk SET aktif=0, diubah=? WHERE uid=?", (now_iso(), uid))
            catat(conn, user, "NONAKTIF_PRODUK", f"{p['kode']} ({n} entri)", ip)
            return {"ok": True, "dinonaktifkan": True, "entri": n}
        conn.execute("DELETE FROM produk WHERE uid=?", (uid,))
        catat(conn, user, "HAPUS_PRODUK", p["kode"], ip)
    return {"ok": True}


# =====================================================================
# PENJUALAN
# =====================================================================
def spg_awal(user):
    with db() as conn:
        toko = rows(conn.execute(
            """SELECT t.uid, t.nama, t.wilayah FROM spg_toko st JOIN toko t ON t.uid=st.toko_uid
               WHERE st.spg_uid=? AND t.aktif=1 ORDER BY t.nama""", (user["uid"],)))
        produk = rows(conn.execute(
            "SELECT uid, kode, nama, harga_botol, harga_dus, isi_dus FROM produk WHERE aktif=1 ORDER BY urut, nama"))
        hari_ini = _entri_list(conn, {"spg_uid": [user["uid"]], "dari": [today_str()], "sampai": [today_str()]})
        terakhir = one(conn.execute(
            """SELECT toko_uid FROM penjualan WHERE spg_uid=? AND dihapus=0 ORDER BY waktu DESC LIMIT 1""", (user["uid"],)))
    return {"toko": toko, "produk": produk, "hari_ini": hari_ini,
            "toko_terakhir": (terakhir or {}).get("toko_uid"), "edit_menit": EDIT_WINDOW_MIN,
            "nama": user["nama"], "jam_server": now_iso()}


def _hitung(produk, satuan, qty, harga, jenis):
    isi = int(produk["isi_dus"] or 1)
    botol = qty * (isi if satuan == "DUS" else 1)
    total = 0 if jenis == "SAMPLE" else round(qty * harga, 2)
    return isi, botol, total


def entri_simpan(user, body, ip, perangkat, uid=None):
    jenis = upper(body.get("jenis") or "JUAL")
    satuan = upper(body.get("satuan") or "BOTOL")
    if jenis not in JENIS:
        raise ApiError(400, "Jenis tidak valid.")
    if satuan not in SATUAN:
        raise ApiError(400, "Satuan tidak valid.")
    qty = angka(body.get("qty"), "Jumlah")
    if qty <= 0:
        raise ApiError(400, "Jumlah harus lebih dari 0.")
    if qty > 100000:
        raise ApiError(400, "Jumlah terlalu besar.")
    harga = angka(body.get("harga"), "Harga")
    if harga < 0:
        raise ApiError(400, "Harga tidak boleh minus.")
    catatan = bersih(body.get("catatan"))[:200]
    with LOCK, db() as conn:
        toko = one(conn.execute("SELECT * FROM toko WHERE uid=? AND aktif=1", (body.get("toko_uid"),)))
        produk = one(conn.execute("SELECT * FROM produk WHERE uid=? AND aktif=1", (body.get("produk_uid"),)))
        if not toko:
            raise ApiError(400, "Toko belum dipilih atau sudah tidak aktif.")
        if not produk:
            raise ApiError(400, "Produk belum dipilih atau sudah tidak aktif.")
        if user["role"] == "spg":
            boleh = conn.execute("SELECT 1 FROM spg_toko WHERE spg_uid=? AND toko_uid=?",
                                 (user["uid"], toko["uid"])).fetchone()
            if not boleh:
                raise ApiError(403, "Toko ini bukan toko Anda. Hubungi master.")
        if jenis == "SAMPLE":
            harga = 0.0
        elif not harga:
            harga = produk["harga_dus"] if satuan == "DUS" else produk["harga_botol"]
        isi, botol, total = _hitung(produk, satuan, qty, harga, jenis)
        if uid:  # edit
            lama = one(conn.execute("SELECT * FROM penjualan WHERE uid=? AND dihapus=0", (uid,)))
            _cek_boleh_ubah(user, lama)
            conn.execute("""UPDATE penjualan SET toko_uid=?, produk_uid=?, jenis=?, satuan=?, qty=?, isi_dus=?, harga=?,
                            total=?, botol_setara=?, catatan=?, diubah=? WHERE uid=?""",
                         (toko["uid"], produk["uid"], jenis, satuan, qty, isi, harga, total, botol, catatan, now_iso(), uid))
            catat(conn, user, "UBAH_PENJUALAN",
                  f"{toko['nama']} | {produk['kode']} | {qty:g} {satuan} | {jenis} | Rp {total:,.0f}", ip, perangkat)
        else:
            uid = new_uid()
            w = now()
            conn.execute("""INSERT INTO penjualan (uid, waktu, tanggal, spg_uid, toko_uid, produk_uid, jenis, satuan,
                            qty, isi_dus, harga, total, botol_setara, catatan, ip, perangkat, dihapus, diubah)
                            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,0,?)""",
                         (uid, w.isoformat(timespec="seconds"), w.date().isoformat(), user["uid"], toko["uid"],
                          produk["uid"], jenis, satuan, qty, isi, harga, total, botol, catatan, ip,
                          (perangkat or "")[:200], now_iso()))
            catat(conn, user, "INPUT_PENJUALAN",
                  f"{toko['nama']} | {produk['kode']} | {qty:g} {satuan} | {jenis} | Rp {total:,.0f}", ip, perangkat)
        hasil = one(conn.execute(_ENTRI_SQL + " WHERE j.uid=?", (uid,)))
    return hasil


def _cek_boleh_ubah(user, entri):
    if not entri:
        raise ApiError(404, "Data tidak ditemukan.")
    if user["role"] == "master":
        return
    if entri["spg_uid"] != user["uid"]:
        raise ApiError(403, "Ini bukan data Anda.")
    batas = datetime.fromisoformat(entri["waktu"]) + timedelta(minutes=EDIT_WINDOW_MIN)
    if now() > batas:
        raise ApiError(403, f"Data yang sudah lewat {EDIT_WINDOW_MIN} menit tidak bisa diubah. Hubungi master.")


def entri_hapus(user, uid, ip):
    with LOCK, db() as conn:
        e = one(conn.execute(_ENTRI_SQL + " WHERE j.uid=?", (uid,)))
        _cek_boleh_ubah(user, e)
        conn.execute("UPDATE penjualan SET dihapus=1, diubah=? WHERE uid=?", (now_iso(), uid))
        catat(conn, user, "HAPUS_PENJUALAN",
              f"{e['toko']} | {e['kode']} | {e['qty']:g} {e['satuan']} | {e['jenis']} | Rp {e['total']:,.0f}", ip)
    return {"ok": True}


_ENTRI_SQL = """
SELECT j.uid, j.waktu, j.tanggal, j.jenis, j.satuan, j.qty, j.isi_dus, j.harga, j.total, j.botol_setara,
       j.catatan, j.spg_uid, j.toko_uid, j.produk_uid, j.dihapus,
       p.nama AS spg, p.username AS spg_user, t.nama AS toko, t.wilayah,
       pr.kode, pr.nama AS produk
FROM penjualan j
JOIN pengguna p ON p.uid=j.spg_uid
JOIN toko t ON t.uid=j.toko_uid
JOIN produk pr ON pr.uid=j.produk_uid
"""


def _filter(q):
    where, args = ["j.dihapus=0"], []
    for key, kolom in (("dari", "j.tanggal >= ?"), ("sampai", "j.tanggal <= ?"), ("spg_uid", "j.spg_uid = ?"),
                       ("toko_uid", "j.toko_uid = ?"), ("produk_uid", "j.produk_uid = ?"),
                       ("jenis", "j.jenis = ?"), ("satuan", "j.satuan = ?"), ("wilayah", "t.wilayah = ?")):
        v = q1(q, key)
        if v:
            where.append(kolom)
            args.append(v)
    if q1(q, "q"):
        where.append("(t.nama LIKE ? OR pr.nama LIKE ? OR pr.kode LIKE ? OR p.nama LIKE ? OR j.catatan LIKE ?)")
        args += [f"%{q1(q, 'q')}%"] * 5
    return " WHERE " + " AND ".join(where), args


def _entri_list(conn, q, limit=2000):
    w, args = _filter(q)
    data = rows(conn.execute(_ENTRI_SQL + w + " ORDER BY j.waktu DESC LIMIT ?", (*args, limit)))
    r = one(conn.execute(
        f"""SELECT COUNT(*) AS jumlah, COALESCE(SUM(j.total),0) AS total,
                   COALESCE(SUM(CASE WHEN j.jenis='JUAL' THEN j.botol_setara ELSE 0 END),0) AS botol_jual,
                   COALESCE(SUM(CASE WHEN j.jenis='SAMPLE' THEN j.botol_setara ELSE 0 END),0) AS botol_sample,
                   COUNT(DISTINCT j.toko_uid) AS toko, COUNT(DISTINCT j.spg_uid) AS spg
            FROM penjualan j JOIN pengguna p ON p.uid=j.spg_uid JOIN toko t ON t.uid=j.toko_uid
            JOIN produk pr ON pr.uid=j.produk_uid {w}""", args))
    return {"rows": data, "ringkas": r}


def entri_list(user, q):
    if user["role"] == "spg":
        q = {**q, "spg_uid": [user["uid"]]}
    with db() as conn:
        return _entri_list(conn, q)


def laporan(q):
    w, args = _filter(q)
    with db() as conn:
        hasil = _entri_list(conn, q, limit=int(q1(q, "limit", "1000")))
        grup = {}
        for nama, kolom in (("spg", "p.nama"), ("toko", "t.nama"), ("produk", "pr.kode || ' - ' || pr.nama"),
                            ("tanggal", "j.tanggal"), ("wilayah", "COALESCE(NULLIF(t.wilayah,''),'(tanpa wilayah)')")):
            grup[nama] = rows(conn.execute(
                f"""SELECT {kolom} AS label,
                           COUNT(*) AS entri,
                           SUM(CASE WHEN j.jenis='JUAL' THEN j.botol_setara ELSE 0 END) AS botol_jual,
                           SUM(CASE WHEN j.jenis='SAMPLE' THEN j.botol_setara ELSE 0 END) AS botol_sample,
                           SUM(j.total) AS total
                    FROM penjualan j JOIN pengguna p ON p.uid=j.spg_uid JOIN toko t ON t.uid=j.toko_uid
                    JOIN produk pr ON pr.uid=j.produk_uid {w}
                    GROUP BY label ORDER BY total DESC, botol_jual DESC""", args))
    hasil["grup"] = grup
    return hasil


def dashboard():
    hari = today_str()
    bulan = hari[:8] + "01"
    with db() as conn:
        f = lambda sql, *a: one(conn.execute(sql, a))
        hari_ini = f("""SELECT COUNT(*) AS entri, COALESCE(SUM(total),0) AS total,
                               COALESCE(SUM(CASE WHEN jenis='JUAL' THEN botol_setara END),0) AS botol,
                               COALESCE(SUM(CASE WHEN jenis='SAMPLE' THEN botol_setara END),0) AS sample,
                               COUNT(DISTINCT spg_uid) AS spg, COUNT(DISTINCT toko_uid) AS toko
                        FROM penjualan WHERE dihapus=0 AND tanggal=?""", hari)
        bulan_ini = f("""SELECT COUNT(*) AS entri, COALESCE(SUM(total),0) AS total,
                                COALESCE(SUM(CASE WHEN jenis='JUAL' THEN botol_setara END),0) AS botol,
                                COALESCE(SUM(CASE WHEN jenis='SAMPLE' THEN botol_setara END),0) AS sample
                         FROM penjualan WHERE dihapus=0 AND tanggal>=?""", bulan)
        terbaru = rows(conn.execute(_ENTRI_SQL + " WHERE j.dihapus=0 ORDER BY j.waktu DESC LIMIT 25"))
        per_spg = rows(conn.execute(
            """SELECT p.nama AS spg, p.uid,
                      COALESCE(SUM(CASE WHEN j.jenis='JUAL' THEN j.botol_setara END),0) AS botol,
                      COALESCE(SUM(CASE WHEN j.jenis='SAMPLE' THEN j.botol_setara END),0) AS sample,
                      COALESCE(SUM(j.total),0) AS total, COUNT(j.uid) AS entri, MAX(j.waktu) AS terakhir
               FROM pengguna p LEFT JOIN penjualan j ON j.spg_uid=p.uid AND j.dihapus=0 AND j.tanggal=?
               WHERE p.role='spg' AND p.aktif=1 GROUP BY p.uid ORDER BY total DESC, p.nama""", (hari,)))
        belum = rows(conn.execute(
            """SELECT t.nama, t.wilayah FROM toko t WHERE t.aktif=1
               AND NOT EXISTS (SELECT 1 FROM penjualan j WHERE j.toko_uid=t.uid AND j.dihapus=0 AND j.tanggal=?)
               ORDER BY t.nama LIMIT 50""", (hari,)))
    return {"hari_ini": hari_ini, "bulan_ini": bulan_ini, "terbaru": terbaru, "per_spg": per_spg,
            "toko_belum_isi": belum, "jam": now_iso()}


def log_list(q):
    where, args = [], []
    if q1(q, "dari"):
        where.append("waktu >= ?"); args.append(q1(q, "dari"))
    if q1(q, "sampai"):
        where.append("waktu < ?"); args.append((date.fromisoformat(q1(q, "sampai")) + timedelta(days=1)).isoformat())
    if q1(q, "aksi"):
        where.append("aksi = ?"); args.append(q1(q, "aksi"))
    if q1(q, "q"):
        where.append("(detail LIKE ? OR username LIKE ? OR ip LIKE ?)"); args += [f"%{q1(q, 'q')}%"] * 3
    w = (" WHERE " + " AND ".join(where)) if where else ""
    with db() as conn:
        return {"rows": rows(conn.execute(f"SELECT * FROM log {w} ORDER BY waktu DESC LIMIT 1000", args)),
                "aksi": [r[0] for r in conn.execute("SELECT DISTINCT aksi FROM log ORDER BY aksi")]}


# =====================================================================
# EXPORT (Excel bila openpyxl ada, kalau tidak CSV)
# =====================================================================
try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    ADA_EXCEL = True
except ImportError:
    ADA_EXCEL = False

KOLOM = ["TANGGAL", "JAM", "SPG", "TOKO", "WILAYAH", "KODE", "PRODUK", "JENIS", "SATUAN", "QTY",
         "ISI/DUS", "BOTOL SETARA", "HARGA", "TOTAL", "CATATAN"]


def _baris_export(r):
    return [r["tanggal"], r["waktu"][11:16], r["spg"], r["toko"], r["wilayah"], r["kode"], r["produk"],
            r["jenis"], r["satuan"], r["qty"], r["isi_dus"], r["botol_setara"], r["harga"], r["total"], r["catatan"]]


def export_penjualan(user, q, ip):
    with db() as conn:
        data = _entri_list(conn, {**q, **({"spg_uid": [user["uid"]]} if user["role"] == "spg" else {})}, limit=100000)
        catat(conn, user, "EXPORT", f"{len(data['rows'])} baris", ip)
    rows_ = list(reversed(data["rows"]))
    nama = f"PENJUALAN_SPG_{now():%Y%m%d_%H%M}"
    if not ADA_EXCEL:
        buf = io.StringIO()
        buf.write(";".join(KOLOM) + "\n")
        for r in rows_:
            buf.write(";".join(str(x).replace(";", ",") for x in _baris_export(r)) + "\n")
        return buf.getvalue().encode("utf-8-sig"), nama + ".csv", "text/csv; charset=utf-8"
    wb = Workbook()
    ws = wb.active
    ws.title = "PENJUALAN"
    ws.append(KOLOM)
    for c in range(1, len(KOLOM) + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3D73")
        cell.alignment = Alignment(horizontal="center")
    for r in rows_:
        ws.append(_baris_export(r))
    n = ws.max_row
    for c in (10, 12, 13, 14):
        for rr in range(2, n + 1):
            ws.cell(rr, c).number_format = "#,##0"
    if n >= 2:
        ws.append([])
        ws.cell(n + 2, 9, "TOTAL")
        ws.cell(n + 2, 10, f"=SUM(J2:J{n})")
        ws.cell(n + 2, 12, f"=SUM(L2:L{n})")
        ws.cell(n + 2, 14, f"=SUM(N2:N{n})")
        for c in (9, 10, 12, 14):
            ws.cell(n + 2, c).font = Font(bold=True)
            ws.cell(n + 2, c).number_format = "#,##0"
    for i, w in enumerate([11, 7, 18, 28, 14, 10, 24, 9, 8, 8, 8, 12, 12, 14, 24], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:O{max(n, 1)}"
    wb.calculation.fullCalcOnLoad = True
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue(), nama + ".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# =====================================================================
# HTTP
# =====================================================================
MASTER, SPG, SEMUA = ("master",), ("spg",), ("master", "spg")
ROUTES = []


def route(method, pattern, roles):
    def deco(fn):
        ROUTES.append((method, re.compile(pattern), roles, fn))
        return fn
    return deco


class Handler(BaseHTTPRequestHandler):
    server_version = "spg"
    sys_version = ""
    protocol_version = "HTTP/1.1"

    def log_message(self, *a):
        pass

    # ---------- util
    @property
    def ip(self):
        if TRUST_PROXY:
            fwd = self.headers.get("X-Forwarded-For", "")
            if fwd:
                return fwd.split(",")[0].strip()[:45]
        return self.client_address[0]

    @property
    def perangkat(self):
        return self.headers.get("User-Agent", "")[:200]

    @property
    def https(self):
        if TRUST_PROXY:
            return self.headers.get("X-Forwarded-Proto", "").lower() == "https"
        return False

    def kirim(self, status, body, ctype, extra=None, nonce=""):
        if isinstance(body, str):
            body = body.encode()
        self.send_response(status)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("X-Frame-Options", "DENY")
        self.send_header("Referrer-Policy", "same-origin")
        self.send_header("Permissions-Policy", "geolocation=(), microphone=(), camera=(), interest-cohort=()")
        self.send_header("Content-Security-Policy",
                         "default-src 'self'; base-uri 'none'; form-action 'self'; frame-ancestors 'none'; "
                         "object-src 'none'; img-src 'self' data:; connect-src 'self'; "
                         f"style-src 'self' 'nonce-{nonce}'; script-src 'self' 'nonce-{nonce}'")
        if self.https:
            self.send_header("Strict-Transport-Security", "max-age=31536000; includeSubDomains")
        for k, v in (extra or {}).items():
            self.send_header(k, v)
        self.end_headers()
        if self.command != "HEAD":
            self.wfile.write(body)

    def json(self, obj, status=200, extra=None):
        self.kirim(status, json.dumps(obj, ensure_ascii=False, default=str), "application/json; charset=utf-8", extra)

    def file(self, data, nama, ctype):
        self.kirim(200, data, ctype, {"Content-Disposition": f'attachment; filename="{nama}"'})

    def raw(self, maks=1_000_000):
        n = int(self.headers.get("Content-Length") or 0)
        if n > maks:
            raise ApiError(413, "Data terlalu besar.")
        return self.rfile.read(n) if n else b""

    def body(self):
        b = self.raw()
        try:
            return json.loads(b.decode()) if b else {}
        except (ValueError, UnicodeDecodeError):
            raise ApiError(400, "Format data tidak valid.")

    def token(self):
        for bagian in (self.headers.get("Cookie") or "").split(";"):
            k, _, v = bagian.strip().partition("=")
            if k == "sid":
                return v
        return ""

    def set_cookie(self, token, hapus=False):
        bagian = [f"sid={'' if hapus else token}", "HttpOnly", "SameSite=Strict", "Path=/"]
        if hapus:
            bagian.append("Max-Age=0")
        if self.https:
            bagian.append("Secure")
        return {"Set-Cookie": "; ".join(bagian)}

    # ---------- routing
    def do_GET(self): self.jalan("GET")
    def do_POST(self): self.jalan("POST")
    def do_PUT(self): self.jalan("PUT")
    def do_DELETE(self): self.jalan("DELETE")

    def jalan(self, method):
        url = urlparse(self.path)
        path, q = url.path.rstrip("/") or "/", parse_qs(url.query)
        try:
            if path == "/healthz":
                return self.kirim(200, "ok", "text/plain")
            # paksa HTTPS di internet
            if REQUIRE_HTTPS and not DEV_MODE and TRUST_PROXY and not self.https:
                host = self.headers.get("Host", "")
                if method == "GET" and host:
                    return self.kirim(308, b"", "text/plain", {"Location": f"https://{host}{self.path}"})
                raise ApiError(400, "Koneksi harus HTTPS.")
            rate_limit(self.ip, "umum", 600, 60)
            user = None
            with db() as conn:
                sesi = ambil_sesi(conn, self.token())
            if sesi:
                user = {"uid": sesi["uid"], "username": sesi["username"], "nama": sesi["nama"],
                        "role": sesi["role"], "csrf": sesi["csrf"], "harus_ganti": sesi["harus_ganti"],
                        "totp_aktif": sesi["totp_aktif"]}
            for m, pat, roles, fn in ROUTES:
                if m != method:
                    continue
                cocok = pat.fullmatch(path)
                if not cocok:
                    continue
                if roles is not None:
                    if not user:
                        raise ApiError(401, "Sesi berakhir. Silakan masuk lagi.")
                    if user["role"] not in roles:
                        raise ApiError(403, "Menu ini tidak tersedia untuk akun Anda.")
                    if method != "GET":
                        if not hmac.compare_digest(self.headers.get("X-CSRF-Token", ""), user["csrf"]):
                            raise ApiError(403, "Permintaan ditolak (token tidak cocok). Muat ulang halaman.")
                        if user["harus_ganti"] and not path.startswith("/api/password"):
                            raise ApiError(403, "Ganti password dulu sebelum memakai aplikasi.")
                return fn(self, q, user, *cocok.groups())
            raise ApiError(404, "Halaman tidak ditemukan.")
        except ApiError as e:
            self.json({"error": e.message, **e.extra}, e.status)
        except (ValueError, KeyError) as e:
            self.json({"error": f"Data tidak valid: {e}"}, 400)
        except BrokenPipeError:
            pass
        except Exception as e:  # noqa
            import traceback
            traceback.print_exc()
            self.json({"error": "Terjadi kesalahan di server."}, 500)


# ---------- halaman
@route("GET", r"/", None)
def r_index(h, q, user):
    nonce = secrets.token_urlsafe(16)
    with db() as conn:
        nama = get_set(conn, "nama_perusahaan", APP_NAME)
    html = INDEX_HTML.replace("__NONCE__", nonce).replace("__APPNAME__", nama)
    h.kirim(200, html, "text/html; charset=utf-8", nonce=nonce)


# ---------- auth
@route("POST", r"/api/login", None)
def r_login(h, q, user):
    token, info = login(h.body(), h.ip, h.perangkat)
    h.json(info, extra=h.set_cookie(token))


@route("POST", r"/api/logout", SEMUA)
def r_logout(h, q, user):
    logout(user, h.token(), h.ip)
    h.json({"ok": True}, extra=h.set_cookie("", hapus=True))


@route("GET", r"/api/me", None)
def r_me(h, q, user):
    if not user:
        return h.json({"login": False})
    h.json({"login": True, **{k: user[k] for k in ("uid", "username", "nama", "role", "csrf", "harus_ganti", "totp_aktif")}})


@route("POST", r"/api/password", SEMUA)
def r_password(h, q, user):
    h.json(ganti_password(user, {**h.body(), "_token": h.token()}, h.ip))


@route("POST", r"/api/2fa/siapkan", SEMUA)
def r_2fa1(h, q, user): h.json(totp_siapkan(user))


@route("POST", r"/api/2fa/aktifkan", SEMUA)
def r_2fa2(h, q, user): h.json(totp_aktifkan(user, h.body(), h.ip))


@route("POST", r"/api/2fa/matikan", SEMUA)
def r_2fa3(h, q, user): h.json(totp_matikan(user, h.body(), h.ip))


@route("GET", r"/api/sesi", SEMUA)
def r_sesi(h, q, user): h.json(sesi_saya(user))


# ---------- spg
@route("GET", r"/api/spg/awal", SEMUA)
def r_spg_awal(h, q, user): h.json(spg_awal(user))


@route("POST", r"/api/penjualan", SEMUA)
def r_jual_baru(h, q, user): h.json(entri_simpan(user, h.body(), h.ip, h.perangkat))


@route("PUT", r"/api/penjualan/(\w+)", SEMUA)
def r_jual_ubah(h, q, user, uid): h.json(entri_simpan(user, h.body(), h.ip, h.perangkat, uid))


@route("DELETE", r"/api/penjualan/(\w+)", SEMUA)
def r_jual_hapus(h, q, user, uid): h.json(entri_hapus(user, uid, h.ip))


@route("GET", r"/api/penjualan", SEMUA)
def r_jual_list(h, q, user): h.json(entri_list(user, q))


@route("GET", r"/export/penjualan", SEMUA)
def r_export(h, q, user):
    data, nama, ctype = export_penjualan(user, q, h.ip)
    h.file(data, nama, ctype)


# ---------- master
@route("GET", r"/api/dashboard", MASTER)
def r_dash(h, q, user): h.json(dashboard())


@route("GET", r"/api/laporan", MASTER)
def r_laporan(h, q, user): h.json(laporan(q))


@route("GET", r"/api/pengguna", MASTER)
def r_user_list(h, q, user): h.json(user_list())


@route("POST", r"/api/pengguna", MASTER)
def r_user_baru(h, q, user): h.json(user_simpan(user, h.body(), None, h.ip))


@route("PUT", r"/api/pengguna/(\w+)", MASTER)
def r_user_ubah(h, q, user, uid): h.json(user_simpan(user, h.body(), uid, h.ip))


@route("DELETE", r"/api/pengguna/(\w+)", MASTER)
def r_user_hapus(h, q, user, uid): h.json(user_hapus(user, uid, h.ip))


@route("PUT", r"/api/pengguna/(\w+)/toko", MASTER)
def r_user_toko(h, q, user, uid): h.json(user_toko(user, uid, h.body(), h.ip))


@route("POST", r"/api/pengguna/(\w+)/reset", MASTER)
def r_user_reset(h, q, user, uid): h.json(user_reset_pw(user, uid, h.ip))


@route("POST", r"/api/pengguna/(\w+)/buka", MASTER)
def r_user_buka(h, q, user, uid): h.json(user_buka_kunci(user, uid, h.ip))


@route("GET", r"/api/toko", SEMUA)
def r_toko_list(h, q, user):
    h.json(toko_list(q) if user["role"] == "master" else spg_awal(user)["toko"])


@route("POST", r"/api/toko", MASTER)
def r_toko_baru(h, q, user): h.json(toko_simpan(user, h.body(), None, h.ip))


@route("PUT", r"/api/toko/(\w+)", MASTER)
def r_toko_ubah(h, q, user, uid): h.json(toko_simpan(user, h.body(), uid, h.ip))


@route("DELETE", r"/api/toko/(\w+)", MASTER)
def r_toko_hapus(h, q, user, uid): h.json(toko_hapus(user, uid, h.ip))


@route("GET", r"/api/produk", SEMUA)
def r_produk(h, q, user): h.json(produk_list(q1(q, "semua") != "1" or user["role"] != "master"))


@route("POST", r"/api/produk", MASTER)
def r_produk_baru(h, q, user): h.json(produk_simpan(user, h.body(), None, h.ip))


@route("PUT", r"/api/produk/(\w+)", MASTER)
def r_produk_ubah(h, q, user, uid): h.json(produk_simpan(user, h.body(), uid, h.ip))


@route("DELETE", r"/api/produk/(\w+)", MASTER)
def r_produk_hapus(h, q, user, uid): h.json(produk_hapus(user, uid, h.ip))


@route("GET", r"/api/log", MASTER)
def r_log(h, q, user): h.json(log_list(q))


INDEX_HTML = r'''<!DOCTYPE html>
<html lang="id">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
<meta name="theme-color" content="#0f766e">
<title>__APPNAME__</title>
<style nonce="__NONCE__">
:root{
  --ink:#16202b; --muted:#5c6b7a; --line:#dfe5ec; --bg:#eef2f6; --card:#fff;
  --brand:#0f766e; --brand-2:#0b5f59; --brand-soft:#d9f2ef;
  --ok:#15803d; --ok-soft:#dcfce7; --warn:#a16207; --warn-soft:#fef3c7;
  --err:#b91c1c; --err-soft:#fee2e2; --row:#f7fafc;
  --pad: env(safe-area-inset-bottom, 0px);
}
*{box-sizing:border-box;-webkit-tap-highlight-color:transparent}
html,body{margin:0;min-height:100%}
body{font-family:"Segoe UI",Roboto,system-ui,Arial,sans-serif;font-size:16px;color:var(--ink);background:var(--bg)}
button,input,select,textarea{font:inherit;color:inherit}
input,select,textarea{border:1px solid #c3ccd6;border-radius:10px;padding:12px;background:#fff;width:100%}
input:focus,select:focus,textarea:focus{outline:3px solid #99e6df;border-color:var(--brand)}
.btn{border:1px solid #c3ccd6;background:#fff;border-radius:10px;padding:12px 16px;cursor:pointer;font-weight:600;
     display:inline-flex;align-items:center;justify-content:center;gap:6px;text-decoration:none;color:inherit}
.btn:active{transform:scale(.98)}
.btn.primary{background:var(--brand);border-color:var(--brand);color:#fff}
.btn.ok{background:var(--ok);border-color:var(--ok);color:#fff}
.btn.danger{color:var(--err);border-color:#e9b4b4;background:#fff}
.btn.block{width:100%}
.btn.sm{padding:6px 10px;font-size:14px;font-weight:500;border-radius:8px}
.btn.big{padding:18px;font-size:19px}
.btn:disabled{opacity:.5}
.card{background:var(--card);border:1px solid var(--line);border-radius:14px;padding:16px;margin-bottom:14px}
.card h2{font-size:15px;margin:0 0 10px;color:var(--muted);text-transform:uppercase;letter-spacing:.04em}
label.f{display:block;font-size:13px;color:var(--muted);margin:0 0 6px}
.muted{color:var(--muted)}
.right{text-align:right}
.center{text-align:center}
.row{display:flex;gap:10px;flex-wrap:wrap;align-items:center}
.row > *{flex:1;min-width:0}
.spacer{flex:1}
.badge{display:inline-block;border-radius:999px;padding:2px 10px;font-size:12px;font-weight:700}
.b-jual{background:var(--brand-soft);color:var(--brand-2)}
.b-sample{background:var(--warn-soft);color:var(--warn)}
.b-ok{background:var(--ok-soft);color:var(--ok)}
.b-err{background:var(--err-soft);color:var(--err)}
.note{padding:12px 14px;border-radius:10px;margin:10px 0;font-size:14px}
.note.ok{background:var(--ok-soft);color:var(--ok)}
.note.warn{background:var(--warn-soft);color:var(--warn)}
.note.err{background:var(--err-soft);color:var(--err)}
.note.info{background:var(--brand-soft);color:var(--brand-2)}
/* ---------- login ---------- */
#login{min-height:100vh;display:flex;align-items:center;justify-content:center;padding:18px;
       background:linear-gradient(160deg,#0b5f59,#134e4a)}
#login .card{width:min(420px,100%);margin:0}
#login h1{font-size:22px;margin:0 0 4px}
#login p.sub{margin:0 0 18px;color:var(--muted);font-size:14px}
/* ---------- app shell ---------- */
#app{display:none;min-height:100vh;flex-direction:column}
header{background:var(--brand);color:#fff;padding:12px 16px;display:flex;align-items:center;gap:12px;position:sticky;top:0;z-index:20}
header b{display:block;font-size:16px;line-height:1.2}
header small{opacity:.85;font-size:12px}
header .btn{background:rgba(255,255,255,.14);border-color:rgba(255,255,255,.3);color:#fff;padding:8px 12px;font-size:14px}
main{flex:1;padding:14px;max-width:1400px;width:100%;margin:0 auto;padding-bottom:calc(24px + var(--pad))}
nav.tabs{display:flex;gap:6px;overflow-x:auto;padding:10px 14px 0;background:var(--brand);position:sticky;top:56px;z-index:19}
nav.tabs button{background:rgba(255,255,255,.12);border:0;color:#e8fffb;padding:10px 14px;border-radius:10px 10px 0 0;cursor:pointer;white-space:nowrap;font-weight:600}
nav.tabs button.active{background:var(--bg);color:var(--brand-2)}
/* ---------- spg ---------- */
.pilihan{display:grid;grid-template-columns:repeat(auto-fill,minmax(150px,1fr));gap:10px}
.pilihan button{background:#fff;border:2px solid var(--line);border-radius:12px;padding:12px;text-align:left;cursor:pointer}
.pilihan button b{display:block;font-size:15px}
.pilihan button small{color:var(--muted)}
.pilihan button.aktif{border-color:var(--brand);background:var(--brand-soft)}
.seg{display:flex;gap:10px}
.seg button{flex:1;padding:14px;border:2px solid var(--line);background:#fff;border-radius:12px;font-weight:700;cursor:pointer}
.seg button.aktif{border-color:var(--brand);background:var(--brand-soft);color:var(--brand-2)}
.qty{display:flex;gap:10px;align-items:center}
.qty button{width:64px;height:56px;font-size:26px;font-weight:700;border-radius:12px;border:2px solid var(--line);background:#fff;cursor:pointer}
.qty input{text-align:center;font-size:24px;font-weight:700;height:56px}
.total-besar{font-size:28px;font-weight:800;color:var(--brand-2);text-align:right;white-space:nowrap}
.baris-total{display:flex;align-items:baseline;justify-content:space-between;gap:10px;flex-wrap:nowrap}
.switch{display:flex;align-items:center;gap:12px;padding:12px;border:2px solid var(--line);border-radius:12px;background:#fff;cursor:pointer}
.switch.aktif{border-color:var(--warn);background:var(--warn-soft)}
.switch input{width:22px;height:22px;flex:none}
.entri{display:flex;gap:10px;align-items:center;padding:12px 0;border-bottom:1px solid var(--line)}
.entri:last-child{border-bottom:0}
.entri .jam{font-variant-numeric:tabular-nums;font-weight:700;width:52px;flex:none}
.entri .isi{flex:1;min-width:0}
.entri .isi b{display:block;font-size:15px;line-height:1.3}
.entri .isi small{color:var(--muted)}
.entri .nilai{font-weight:700;white-space:nowrap;font-size:15px}
.ring{display:grid;grid-template-columns:repeat(auto-fit,minmax(110px,1fr));gap:10px;margin-bottom:12px}
.ring div{background:var(--card);border:1px solid var(--line);border-radius:12px;padding:10px 12px}
.ring b{display:block;font-size:19px;font-variant-numeric:tabular-nums;overflow-wrap:anywhere;line-height:1.25}
.ring span{font-size:12px;color:var(--muted)}
/* ---------- daftar cari ---------- */
.cari{position:relative}
.cari .hasil{position:absolute;left:0;right:0;top:100%;background:#fff;border:1px solid #c3ccd6;border-radius:10px;
              box-shadow:0 12px 30px rgba(0,0,0,.18);max-height:280px;overflow:auto;z-index:40;display:none;margin-top:4px}
.cari .hasil div{padding:14px 12px;cursor:pointer;border-bottom:1px solid #f0f3f6}
.cari .hasil div small{color:var(--muted);display:block}
.cari .hasil div.hl,.cari .hasil div:hover{background:var(--brand-soft)}
.terpilih{display:flex;align-items:center;gap:10px;background:var(--brand-soft);border:2px solid var(--brand);
          border-radius:12px;padding:14px}
.terpilih b{flex:1;font-size:17px}
/* ---------- tabel ---------- */
.tbl{overflow:auto;border:1px solid var(--line);border-radius:12px;background:#fff;max-height:70vh}
table{border-collapse:collapse;width:100%;font-size:14px}
th,td{padding:8px 10px;border-bottom:1px solid var(--line);white-space:nowrap;text-align:left}
th{background:var(--brand);color:#fff;position:sticky;top:0;font-size:12px;text-transform:uppercase}
tbody tr:nth-child(even){background:var(--row)}
td.n{text-align:right;font-variant-numeric:tabular-nums}
tfoot td{position:sticky;bottom:0;background:var(--brand-soft);font-weight:700}
.empty{padding:28px;text-align:center;color:var(--muted)}
.grid2{display:grid;grid-template-columns:repeat(auto-fit,minmax(320px,1fr));gap:14px}
.form{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:12px}
.form .full{grid-column:1/-1}
.chips{display:flex;flex-wrap:wrap;gap:6px}
.chip{background:var(--brand-soft);color:var(--brand-2);border-radius:999px;padding:3px 10px;font-size:12px}
.checklist{max-height:300px;overflow:auto;border:1px solid var(--line);border-radius:10px;padding:8px}
.checklist label{display:flex;gap:8px;align-items:center;padding:8px;border-bottom:1px solid #f1f4f7}
.checklist input{width:20px;height:20px;flex:none}
code{background:#eef2f7;padding:2px 6px;border-radius:6px;font-size:14px}
dialog{border:0;border-radius:14px;padding:0;width:min(640px,94vw);box-shadow:0 20px 60px rgba(0,0,0,.3)}
dialog::backdrop{background:rgba(10,20,30,.55)}
dialog .h{padding:14px 16px;border-bottom:1px solid var(--line);display:flex;align-items:center;gap:10px}
dialog .h h3{margin:0;font-size:17px;flex:1}
dialog .b{padding:16px;max-height:70vh;overflow:auto}
dialog .f{padding:12px 16px;border-top:1px solid var(--line);display:flex;gap:8px;justify-content:flex-end}
#toast{position:fixed;left:50%;bottom:calc(20px + var(--pad));transform:translateX(-50%);background:#16202b;color:#fff;
       padding:14px 20px;border-radius:12px;display:none;z-index:100;max-width:92vw;font-weight:600;text-align:center}
#toast.err{background:var(--err)}
#toast.ok{background:var(--ok)}
@media (max-width:640px){
  main{padding:10px}
  .ring{grid-template-columns:repeat(2,1fr)}
  .ring b{font-size:17px}
  .total-besar{font-size:24px}
  .entri .isi b{font-size:14px}
  .entri .isi small{font-size:12px}
  .card{padding:13px;border-radius:12px}
  nav.tabs{top:52px}
}
</style>
</head>
<body>

<section id="login">
  <div class="card">
    <h1>__APPNAME__</h1>
    <p class="sub">Masuk untuk mulai mencatat penjualan</p>
    <label class="f" for="u">Username</label>
    <input id="u" autocomplete="username" autocapitalize="none" spellcheck="false">
    <div style="height:12px"></div>
    <label class="f" for="p">Password</label>
    <input id="p" type="password" autocomplete="current-password">
    <div id="kode-wrap" style="display:none">
      <div style="height:12px"></div>
      <label class="f" for="k">Kode 6 angka (aplikasi Authenticator)</label>
      <input id="k" inputmode="numeric" maxlength="6" autocomplete="one-time-code">
    </div>
    <div id="login-err"></div>
    <div style="height:14px"></div>
    <button class="btn primary big block" id="masuk">Masuk</button>
  </div>
</section>

<section id="app">
  <header>
    <div style="flex:1;min-width:0"><b id="h-nama">-</b><small id="h-info">-</small></div>
    <button class="btn" id="h-akun">Akun</button>
    <button class="btn" id="h-keluar">Keluar</button>
  </header>
  <nav class="tabs" id="tabs" style="display:none"></nav>
  <main id="isi"></main>
</section>

<dialog id="modal">
  <div class="h"><h3 id="m-judul"></h3><button class="btn sm" id="m-tutup">Tutup</button></div>
  <div class="b" id="m-isi"></div>
  <div class="f" id="m-aksi"></div>
</dialog>
<div id="toast" role="status" aria-live="polite"></div>

<script nonce="__NONCE__">
// ============================================================ util
const $ = (s, el = document) => el.querySelector(s);
const $$ = (s, el = document) => [...el.querySelectorAll(s)];
const esc = s => String(s ?? '').replace(/[&<>"']/g, c => ({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));
const num = n => Number(n || 0).toLocaleString('id-ID', {maximumFractionDigits: 2});
const rp = n => 'Rp ' + Math.round(n || 0).toLocaleString('id-ID');
const BLN = ['Jan','Feb','Mar','Apr','Mei','Jun','Jul','Agu','Sep','Okt','Nov','Des'];
const tgl = s => { if (!s) return '-'; const [y, m, d] = s.slice(0, 10).split('-'); return `${+d} ${BLN[m - 1]} ${y}`; };
const jam = s => s ? s.slice(11, 16) : '';
const qs = o => new URLSearchParams(Object.entries(o).filter(([, v]) => v !== '' && v != null)).toString();
let ME = null, TIMER = null;

async function api(url, opt = {}) {
  const init = {method: opt.method || 'GET', headers: {}, credentials: 'same-origin'};
  if (opt.body !== undefined) { init.body = JSON.stringify(opt.body); init.headers['Content-Type'] = 'application/json'; }
  if (init.method !== 'GET') init.headers['X-CSRF-Token'] = ME?.csrf || '';
  const res = await fetch(url, init);
  let data = null; try { data = await res.json(); } catch {}
  if (res.status === 401 && ME) { ME = null; tampilLogin('Sesi berakhir, silakan masuk lagi.'); throw new Error('Sesi berakhir.'); }
  if (!res.ok) { const e = new Error(data?.error || `Gagal (${res.status})`); e.status = res.status; e.data = data || {}; throw e; }
  return data;
}
let tOut;
function toast(pesan, tipe = 'ok') {
  const t = $('#toast'); t.textContent = pesan; t.className = tipe; t.style.display = 'block';
  clearTimeout(tOut); tOut = setTimeout(() => t.style.display = 'none', tipe === 'err' ? 6000 : 2600);
}
const jaga = fn => async (...a) => { try { return await fn(...a); } catch (e) { toast(e.message, 'err'); } };
function unduh(url) { const a = document.createElement('a'); a.href = url; document.body.append(a); a.click(); a.remove(); }
const hariIni = () => { const d = new Date(); return new Date(d - d.getTimezoneOffset() * 6e4).toISOString().slice(0, 10); };
const awalBulan = () => hariIni().slice(0, 8) + '01';

// ============================================================ modal
const modal = $('#modal');
function buka(judul, isi, aksi = []) {
  $('#m-judul').textContent = judul;
  const b = $('#m-isi'); b.innerHTML = '';
  typeof isi === 'string' ? (b.innerHTML = isi) : b.append(isi);
  const f = $('#m-aksi'); f.innerHTML = '';
  aksi.forEach(a => {
    const el = document.createElement('button');
    el.className = 'btn ' + (a.cls || ''); el.textContent = a.teks;
    el.onclick = jaga(async () => { if (await a.klik?.() !== false) tutup(); });
    f.append(el);
  });
  f.style.display = aksi.length ? '' : 'none';
  if (!modal.open) modal.showModal();
}
const tutup = () => modal.open && modal.close();
$('#m-tutup').onclick = tutup;
function konfirmasi(judul, teks, ok = 'Ya', bahaya = true) {
  return new Promise(res => {
    let selesai = false;
    buka(judul, `<p>${teks}</p>`, [
      {teks: 'Batal', klik: () => { selesai = true; res(false); }},
      {teks: ok, cls: bahaya ? 'danger' : 'primary', klik: () => { selesai = true; res(true); }}]);
    modal.addEventListener('close', () => { if (!selesai) res(false); }, {once: true});
  });
}

// ============================================================ login
function tampilLogin(pesan) {
  clearInterval(TIMER);
  $('#app').style.display = 'none'; $('#login').style.display = 'flex';
  $('#login-err').innerHTML = pesan ? `<div class="note err">${esc(pesan)}</div>` : '';
  $('#p').value = ''; $('#k').value = ''; $('#kode-wrap').style.display = 'none';
  setTimeout(() => ($('#u').value ? $('#p') : $('#u')).focus(), 60);
}
const masuk = async () => {
  const btn = $('#masuk'); btn.disabled = true; $('#login-err').innerHTML = '';
  try {
    ME = await api('/api/login', {method: 'POST', body: {username: $('#u').value, password: $('#p').value, kode: $('#k').value}});
    await mulai();
  } catch (e) {
    if (e.data?.butuh_kode) { $('#kode-wrap').style.display = 'block'; setTimeout(() => $('#k').focus(), 50); }
    $('#login-err').innerHTML = `<div class="note err">${esc(e.message)}</div>`;
  } finally { btn.disabled = false; }
};
$('#masuk').onclick = masuk;
['u', 'p', 'k'].forEach(id => $('#' + id).addEventListener('keydown', e => e.key === 'Enter' && masuk()));
$('#h-keluar').onclick = jaga(async () => {
  if (!await konfirmasi('Keluar', 'Keluar dari aplikasi?', 'Keluar')) return;
  await api('/api/logout', {method: 'POST', body: {}}).catch(() => {});
  ME = null; tampilLogin();
});
$('#h-akun').onclick = () => halamanAkun();

async function mulai() {
  $('#login').style.display = 'none'; $('#app').style.display = 'flex';
  $('#h-nama').textContent = ME.nama;
  $('#h-info').textContent = (ME.role === 'master' ? 'Master' : 'SPG') + ' · ' + ME.username;
  if (ME.harus_ganti) return gantiPassword(true);
  ME.role === 'master' ? masterUI() : spgUI();
}

// ============================================================ ganti password / akun
function gantiPassword(wajib) {
  const w = document.createElement('div');
  w.innerHTML = `${wajib ? '<div class="note warn">Demi keamanan, ganti password bawaan Anda sebelum memakai aplikasi.</div>' : ''}
    <label class="f">Password sekarang</label><input type="password" id="pw-lama" autocomplete="current-password">
    <div style="height:10px"></div><label class="f">Password baru</label><input type="password" id="pw-baru" autocomplete="new-password">
    <div style="height:10px"></div><label class="f">Ulangi password baru</label><input type="password" id="pw-baru2" autocomplete="new-password">
    <p class="muted" style="font-size:13px">Minimal 10 karakter, ada huruf dan angka, bukan nama/username Anda.</p>`;
  buka('Ganti Password', w, [
    ...(wajib ? [] : [{teks: 'Batal'}]),
    {teks: 'Simpan', cls: 'primary', klik: async () => {
      if ($('#pw-baru', w).value !== $('#pw-baru2', w).value) { toast('Password baru tidak sama', 'err'); return false; }
      await api('/api/password', {method: 'POST', body: {lama: $('#pw-lama', w).value, baru: $('#pw-baru', w).value}});
      toast('Password berhasil diganti');
      ME.harus_ganti = false;
      ME.role === 'master' ? masterUI() : spgUI();
    }}]);
  if (wajib) modal.addEventListener('cancel', e => e.preventDefault(), {once: true});
}

function halamanAkun() {
  const w = document.createElement('div');
  w.innerHTML = `<div class="row"><button class="btn block" id="a-pw">Ganti password</button></div>
    <div style="height:12px"></div>
    <div class="note ${ME.totp_aktif ? 'ok' : 'warn'}">Verifikasi 2 langkah (kode dari aplikasi Authenticator):
      <b>${ME.totp_aktif ? 'AKTIF' : 'belum aktif'}</b>.
      ${ME.role === 'master' && !ME.totp_aktif ? ' Sangat disarankan diaktifkan untuk akun master.' : ''}</div>
    <button class="btn block ${ME.totp_aktif ? 'danger' : 'primary'}" id="a-2fa">${ME.totp_aktif ? 'Matikan verifikasi 2 langkah' : 'Aktifkan verifikasi 2 langkah'}</button>
    <h3 style="margin:18px 0 8px;font-size:15px">Perangkat yang sedang masuk</h3><div id="a-sesi" class="muted">memuat…</div>`;
  buka('Akun Saya', w, []);
  $('#a-pw', w).onclick = () => gantiPassword(false);
  $('#a-2fa', w).onclick = () => ME.totp_aktif ? matikan2fa() : aktifkan2fa();
  api('/api/sesi').then(list => {
    $('#a-sesi', w).innerHTML = list.map(s => `<div class="entri"><div class="isi"><b>${esc(s.ip || '-')}</b>
      <small>${esc((s.perangkat || '').slice(0, 70))}</small></div><div class="muted">${tgl(s.terakhir)} ${jam(s.terakhir)}</div></div>`).join('');
  }).catch(() => {});
}
const aktifkan2fa = jaga(async () => {
  const d = await api('/api/2fa/siapkan', {method: 'POST', body: {}});
  const w = document.createElement('div');
  w.innerHTML = `<ol style="padding-left:20px;line-height:1.8">
      <li>Buka aplikasi <b>Google Authenticator</b> (atau sejenis) di HP.</li>
      <li>Pilih tambah akun → <b>Masukkan kunci setup</b>.</li>
      <li>Nama akun: <code>${esc(ME.username)}</code></li>
      <li>Kunci: <code style="font-size:16px;letter-spacing:2px">${esc(d.secret)}</code></li>
      <li>Masukkan 6 angka yang muncul di bawah ini.</li></ol>
    <input id="kode2fa" inputmode="numeric" maxlength="6" placeholder="123456">`;
  buka('Aktifkan Verifikasi 2 Langkah', w, [{teks: 'Batal'}, {teks: 'Aktifkan', cls: 'primary', klik: async () => {
    await api('/api/2fa/aktifkan', {method: 'POST', body: {kode: $('#kode2fa', w).value}});
    ME.totp_aktif = true; toast('Verifikasi 2 langkah aktif');
  }}]);
});
const matikan2fa = () => {
  const w = document.createElement('div');
  w.innerHTML = `<div class="note warn">Keamanan akun akan berkurang.</div>
    <label class="f">Masukkan password Anda</label><input type="password" id="pw2fa">`;
  buka('Matikan Verifikasi 2 Langkah', w, [{teks: 'Batal'}, {teks: 'Matikan', cls: 'danger', klik: async () => {
    await api('/api/2fa/matikan', {method: 'POST', body: {password: $('#pw2fa', w).value}});
    ME.totp_aktif = false; toast('Verifikasi 2 langkah dimatikan');
  }}]);
};

// ============================================================ pencarian (toko / produk)
function pencarian(kotak, daftar, {onPilih, teks = 'Ketik nama…', kunci = t => t.nama, sub = () => ''} = {}) {
  kotak.classList.add('cari');
  kotak.innerHTML = `<input placeholder="${teks}" autocomplete="off" inputmode="search"><div class="hasil"></div>`;
  const inp = $('input', kotak), box = $('.hasil', kotak);
  let hasil = [], hl = 0;
  const gambar = () => {
    const q = inp.value.trim().toUpperCase().split(/\s+/).filter(Boolean);
    hasil = daftar().filter(t => q.every(w => (kunci(t) + ' ' + sub(t)).toUpperCase().includes(w))).slice(0, 40);
    hl = 0;
    box.innerHTML = hasil.length ? hasil.map((t, i) => `<div data-i="${i}" class="${i ? '' : 'hl'}">${esc(kunci(t))}
      ${sub(t) ? `<small>${esc(sub(t))}</small>` : ''}</div>`).join('') : '<div class="muted" style="padding:14px">Tidak ditemukan</div>';
    box.style.display = 'block';
  };
  inp.addEventListener('focus', gambar);
  inp.addEventListener('input', gambar);
  inp.addEventListener('blur', () => setTimeout(() => box.style.display = 'none', 180));
  inp.addEventListener('keydown', e => {
    if (box.style.display !== 'block') return;
    const d = $$('[data-i]', box);
    if (e.key === 'ArrowDown') { hl = Math.min(hl + 1, d.length - 1); e.preventDefault(); }
    else if (e.key === 'ArrowUp') { hl = Math.max(hl - 1, 0); e.preventDefault(); }
    else if (e.key === 'Enter') { if (hasil[hl]) { onPilih(hasil[hl]); box.style.display = 'none'; e.preventDefault(); } return; }
    else return;
    d.forEach((x, i) => x.classList.toggle('hl', i === hl)); d[hl]?.scrollIntoView({block: 'nearest'});
  });
  box.addEventListener('mousedown', e => { const d = e.target.closest('[data-i]'); if (d) { onPilih(hasil[+d.dataset.i]); box.style.display = 'none'; } });
  return {fokus: () => inp.focus(), kosong: () => { inp.value = ''; }};
}

// ============================================================ HALAMAN SPG
async function spgUI() {
  $('#tabs').style.display = 'none';
  const isi = $('#isi');
  isi.innerHTML = '<div class="card">Memuat…</div>';
  const d = await api('/api/spg/awal');
  let toko = d.toko.find(t => t.uid === d.toko_terakhir) || null;
  let produk = null, satuan = 'BOTOL', sample = false, qty = 1, hargaManual = null;

  isi.innerHTML = `
    <div class="ring" id="s-ring"></div>
    <div class="card">
      <h2>1. Toko</h2>
      <div id="s-toko"></div>
      ${d.toko.length ? '' : '<div class="note warn">Belum ada toko yang ditugaskan untuk Anda. Hubungi master.</div>'}
    </div>
    <div class="card">
      <h2>2. Produk</h2>
      <div id="s-cariproduk" style="margin-bottom:10px;${d.produk.length > 8 ? '' : 'display:none'}"></div>
      <div class="pilihan" id="s-produk"></div>
    </div>
    <div class="card">
      <h2>3. Jumlah</h2>
      <div class="seg" id="s-satuan">
        <button data-s="BOTOL" class="aktif">BOTOL</button>
        <button data-s="DUS">DUS</button>
      </div>
      <div style="height:12px"></div>
      <div class="qty">
        <button id="s-kurang" aria-label="Kurangi">−</button>
        <input id="s-qty" inputmode="decimal" value="1">
        <button id="s-tambah" aria-label="Tambah">+</button>
      </div>
      <div style="height:12px"></div>
      <div class="switch" id="s-sample"><input type="checkbox" id="s-sample-cb"><div><b>Sample gratis</b><br><small class="muted">harga otomatis jadi 0</small></div></div>
      <div style="height:12px"></div>
      <div id="s-harga-wrap">
        <label class="f">Harga per <span id="s-satuan-label">botol</span> (boleh diubah)</label>
        <input id="s-harga" inputmode="numeric">
      </div>
      <div style="height:12px"></div>
      <label class="f">Catatan (boleh dikosongkan)</label>
      <input id="s-catatan" placeholder="contoh: promo, minta nota">
      <div style="height:14px"></div>
      <div class="baris-total"><span class="muted">TOTAL</span><div class="total-besar" id="s-total">Rp 0</div></div>
      <div style="height:12px"></div>
      <button class="btn ok big block" id="s-kirim">KIRIM</button>
    </div>
    <div class="card">
      <h2>Sudah dikirim hari ini</h2>
      <div id="s-daftar"></div>
    </div>`;

  // --- toko
  const kotakToko = $('#s-toko');
  const gambarToko = () => {
    if (toko) {
      kotakToko.className = '';
      kotakToko.innerHTML = `<div class="terpilih"><b>${esc(toko.nama)}</b><button class="btn sm" id="s-ganti-toko">Ganti</button></div>`;
      $('#s-ganti-toko').onclick = () => { toko = null; gambarToko(); };
    } else {
      kotakToko.innerHTML = '';
      const p = pencarian(kotakToko, () => d.toko, {teks: 'Ketik nama toko…', sub: t => t.wilayah || '',
        onPilih: t => { toko = t; gambarToko(); hitung(); }});
      p.fokus();
    }
    hitung();
  };

  // --- produk
  let filterProduk = '';
  const gambarProduk = () => {
    const list = d.produk.filter(p => !filterProduk || (p.nama + ' ' + p.kode).toUpperCase().includes(filterProduk));
    $('#s-produk').innerHTML = list.map(p => `<button data-p="${p.uid}" class="${produk && produk.uid === p.uid ? 'aktif' : ''}">
      <b>${esc(p.nama)}</b><small>${esc(p.kode)} · ${rp(p.harga_botol)}/botol</small></button>`).join('')
      || '<div class="muted">Produk tidak ada.</div>';
  };
  pencarian($('#s-cariproduk'), () => d.produk, {teks: 'Cari produk…', kunci: p => p.nama, sub: p => p.kode,
    onPilih: p => { pilihProduk(p.uid); }});
  const pilihProduk = uid => {
    produk = d.produk.find(p => p.uid === uid) || null;
    hargaManual = null; gambarProduk(); isiHarga(); hitung();
  };
  $('#s-produk').onclick = e => { const b = e.target.closest('[data-p]'); if (b) pilihProduk(b.dataset.p); };
  gambarProduk();

  // --- satuan, qty, harga
  const hargaDefault = () => !produk ? 0 : (satuan === 'DUS' ? produk.harga_dus : produk.harga_botol);
  const isiHarga = () => { $('#s-harga').value = hargaDefault() ? Math.round(hargaDefault()) : ''; };
  const hitung = () => {
    const h = sample ? 0 : (parseFloat($('#s-harga').value) || 0);
    $('#s-total').textContent = rp(qty * h);
    $('#s-satuan-label').textContent = satuan.toLowerCase();
    $('#s-harga-wrap').style.display = sample ? 'none' : '';
    $('#s-sample').classList.toggle('aktif', sample);
    $('#s-kirim').textContent = sample ? 'KIRIM SAMPLE GRATIS' : 'KIRIM';
    $('#s-kirim').className = 'btn big block ' + (sample ? 'primary' : 'ok');
  };
  $('#s-satuan').onclick = e => {
    const b = e.target.closest('[data-s]'); if (!b) return;
    satuan = b.dataset.s;
    $$('#s-satuan button').forEach(x => x.classList.toggle('aktif', x === b));
    isiHarga(); hitung();
  };
  const setQty = v => { qty = Math.max(0, Math.round(v * 100) / 100); $('#s-qty').value = qty; hitung(); };
  $('#s-kurang').onclick = () => setQty(qty - 1);
  $('#s-tambah').onclick = () => setQty(qty + 1);
  $('#s-qty').oninput = () => { qty = parseFloat($('#s-qty').value) || 0; hitung(); };
  $('#s-harga').oninput = hitung;
  $('#s-sample').onclick = e => {
    if (e.target.id !== 's-sample-cb') $('#s-sample-cb').checked = !$('#s-sample-cb').checked;
    sample = $('#s-sample-cb').checked; hitung();
  };
  gambarToko();

  // --- daftar hari ini
  const gambarDaftar = data => {
    const r = data.ringkas;
    $('#s-ring').innerHTML = `
      <div><b>${num(r.botol_jual)}</b><span>botol terjual</span></div>
      <div><b>${num(r.botol_sample)}</b><span>botol sample</span></div>
      <div><b>${rp(r.total)}</b><span>total rupiah</span></div>
      <div><b>${num(r.jumlah)}</b><span>entri hari ini</span></div>`;
    $('#s-daftar').innerHTML = data.rows.length ? data.rows.map(e => {
      const bolehHapus = (Date.now() - new Date(e.waktu).getTime()) / 60000 < d.edit_menit;
      return `<div class="entri"><div class="jam">${jam(e.waktu)}</div>
        <div class="isi"><b>${esc(e.produk)} · ${num(e.qty)} ${e.satuan}</b>
          <small>${esc(e.toko)} ${e.jenis === 'SAMPLE' ? '<span class="badge b-sample">SAMPLE</span>' : ''}${e.catatan ? ' · ' + esc(e.catatan) : ''}</small></div>
        <div class="nilai">${e.jenis === 'SAMPLE' ? '-' : rp(e.total)}</div>
        ${bolehHapus ? `<button class="btn sm danger" data-hapus="${e.uid}">Hapus</button>` : ''}</div>`;
    }).join('') : '<div class="empty">Belum ada. Isi form di atas lalu tekan KIRIM.</div>';
  };
  gambarDaftar(d.hari_ini);
  $('#s-daftar').onclick = jaga(async e => {
    const b = e.target.closest('[data-hapus]'); if (!b) return;
    if (!await konfirmasi('Hapus', 'Hapus data ini?', 'Hapus')) return;
    await api('/api/penjualan/' + b.dataset.hapus, {method: 'DELETE'});
    toast('Data dihapus'); segarkan();
  });
  const segarkan = jaga(async () => {
    const baru = await api('/api/penjualan?' + qs({dari: hariIni(), sampai: hariIni()}));
    gambarDaftar(baru);
  });

  // --- kirim
  $('#s-kirim').onclick = jaga(async () => {
    if (!toko) return toast('Pilih toko dulu', 'err');
    if (!produk) return toast('Pilih produk dulu', 'err');
    if (!qty) return toast('Isi jumlahnya dulu', 'err');
    const btn = $('#s-kirim'); btn.disabled = true;
    try {
      await api('/api/penjualan', {method: 'POST', body: {
        toko_uid: toko.uid, produk_uid: produk.uid, jenis: sample ? 'SAMPLE' : 'JUAL', satuan,
        qty, harga: sample ? 0 : $('#s-harga').value, catatan: $('#s-catatan').value}});
      toast(sample ? 'Sample tercatat' : 'Penjualan tercatat', 'ok');
      $('#s-catatan').value = ''; setQty(1);
      $('#s-sample-cb').checked = false; sample = false; isiHarga(); hitung();
      segarkan();
    } finally { btn.disabled = false; }
  });
  clearInterval(TIMER);
  TIMER = setInterval(segarkan, 60000);
}

// ============================================================ HALAMAN MASTER
const TAB = [['dashboard', 'Dashboard'], ['penjualan', 'Penjualan'], ['laporan', 'Laporan'],
             ['spg', 'Akun SPG'], ['toko', 'Toko'], ['produk', 'Produk'], ['log', 'Log Keamanan']];
let tabAktif = 'dashboard';
function masterUI() {
  const nav = $('#tabs');
  nav.style.display = 'flex';
  nav.innerHTML = TAB.map(([k, l]) => `<button data-t="${k}" class="${k === tabAktif ? 'active' : ''}">${l}</button>`).join('');
  nav.onclick = e => { const b = e.target.closest('[data-t]'); if (b) gantiTab(b.dataset.t); };
  gantiTab(tabAktif);
}
function gantiTab(t) {
  tabAktif = t;
  $$('#tabs button').forEach(b => b.classList.toggle('active', b.dataset.t === t));
  clearInterval(TIMER);
  $('#isi').innerHTML = '<div class="card">Memuat…</div>';
  jaga({dashboard: mDashboard, penjualan: mPenjualan, laporan: mLaporan, spg: mSpg, toko: mToko,
        produk: mProduk, log: mLog}[t])($('#isi'));
}

// ---------- dashboard realtime
async function mDashboard(el) {
  const gambar = d => {
    const h = d.hari_ini, b = d.bulan_ini;
    el.innerHTML = `
      <div class="ring">
        <div><b>${rp(h.total)}</b><span>penjualan hari ini</span></div>
        <div><b>${num(h.botol)}</b><span>botol terjual hari ini</span></div>
        <div><b>${num(h.sample)}</b><span>botol sample hari ini</span></div>
        <div><b>${num(h.spg)}</b><span>SPG aktif input</span></div>
        <div><b>${num(h.toko)}</b><span>toko terisi</span></div>
        <div><b>${rp(b.total)}</b><span>penjualan bulan ini</span></div>
      </div>
      <div class="grid2">
        <div class="card"><h2>Masuk terbaru (otomatis diperbarui)</h2><div id="d-live"></div></div>
        <div>
          <div class="card"><h2>Per SPG hari ini</h2><div class="tbl" style="max-height:320px"><table>
            <thead><tr><th>SPG</th><th>Entri</th><th>Botol</th><th>Sample</th><th>Total</th><th>Terakhir</th></tr></thead>
            <tbody>${d.per_spg.map(s => `<tr><td>${esc(s.spg)}</td><td class="n">${s.entri}</td><td class="n">${num(s.botol)}</td>
              <td class="n">${num(s.sample)}</td><td class="n">${num(s.total)}</td><td>${s.terakhir ? jam(s.terakhir) : '-'}</td></tr>`).join('')}</tbody>
          </table></div></div>
          <div class="card"><h2>Toko belum ada input hari ini (${d.toko_belum_isi.length})</h2>
            <div class="chips">${d.toko_belum_isi.map(t => `<span class="chip">${esc(t.nama)}</span>`).join('') || '<span class="muted">Semua toko sudah ada input.</span>'}</div></div>
        </div>
      </div>`;
    $('#d-live').innerHTML = d.terbaru.length ? d.terbaru.map(e => `<div class="entri">
      <div class="jam">${jam(e.waktu)}</div><div class="isi"><b>${esc(e.produk)} · ${num(e.qty)} ${e.satuan}</b>
      <small>${esc(e.toko)} · ${esc(e.spg)} ${e.jenis === 'SAMPLE' ? '<span class="badge b-sample">SAMPLE</span>' : ''}</small></div>
      <div class="nilai">${e.jenis === 'SAMPLE' ? '-' : rp(e.total)}</div></div>`).join('') : '<div class="empty">Belum ada input hari ini.</div>';
  };
  gambar(await api('/api/dashboard'));
  TIMER = setInterval(jaga(async () => { if (tabAktif === 'dashboard') gambar(await api('/api/dashboard')); }), 20000);
}

// ---------- penjualan (filter + tabel + export)
function filterBar(id, opsi, onTampil) {
  return `<div class="card"><div class="form">
    <div><label class="f">Dari</label><input type="date" id="${id}-dari" value="${awalBulan()}"></div>
    <div><label class="f">Sampai</label><input type="date" id="${id}-sampai" value="${hariIni()}"></div>
    <div><label class="f">SPG</label><select id="${id}-spg"><option value="">Semua</option>${opsi.spg}</select></div>
    <div><label class="f">Toko</label><select id="${id}-toko"><option value="">Semua</option>${opsi.toko}</select></div>
    <div><label class="f">Produk</label><select id="${id}-produk"><option value="">Semua</option>${opsi.produk}</select></div>
    <div><label class="f">Jenis</label><select id="${id}-jenis"><option value="">Semua</option><option value="JUAL">Jual</option><option value="SAMPLE">Sample</option></select></div>
    <div><label class="f">Cari</label><input id="${id}-q" placeholder="toko / produk / catatan"></div>
    <div class="full row"><button class="btn primary" id="${id}-go">Tampilkan</button>
      <button class="btn" data-cepat="hari">Hari ini</button>
      <button class="btn" data-cepat="kemarin">Kemarin</button>
      <button class="btn" data-cepat="bulan">Bulan ini</button>
      <span class="spacer"></span>${onTampil || ''}</div>
  </div></div>`;
}
const nilaiFilter = id => ({dari: $(`#${id}-dari`).value, sampai: $(`#${id}-sampai`).value, spg_uid: $(`#${id}-spg`).value,
  toko_uid: $(`#${id}-toko`).value, produk_uid: $(`#${id}-produk`).value, jenis: $(`#${id}-jenis`).value, q: $(`#${id}-q`).value});
function pasangCepat(el, id, load) {
  el.addEventListener('click', e => {
    const b = e.target.closest('[data-cepat]'); if (!b) return;
    const d = new Date(), pad = n => String(n).padStart(2, '0');
    const f = x => `${x.getFullYear()}-${pad(x.getMonth() + 1)}-${pad(x.getDate())}`;
    let a, c;
    if (b.dataset.cepat === 'hari') a = c = f(d);
    else if (b.dataset.cepat === 'kemarin') { const y = new Date(d - 864e5); a = c = f(y); }
    else { a = awalBulan(); c = hariIni(); }
    $(`#${id}-dari`).value = a; $(`#${id}-sampai`).value = c; load();
  });
}
async function opsiFilter() {
  const [spg, toko, produk] = await Promise.all([api('/api/pengguna'), api('/api/toko'), api('/api/produk?semua=1')]);
  return {
    spg: spg.filter(u => u.role === 'spg').map(u => `<option value="${u.uid}">${esc(u.nama)}</option>`).join(''),
    toko: toko.map(t => `<option value="${t.uid}">${esc(t.nama)}</option>`).join(''),
    produk: produk.map(p => `<option value="${p.uid}">${esc(p.nama)}</option>`).join(''),
    _spg: spg, _toko: toko, _produk: produk};
}
async function mPenjualan(el) {
  const opsi = await opsiFilter();
  el.innerHTML = filterBar('f', opsi, '<button class="btn" id="f-xls">Export Excel</button>') +
    '<div class="ring" id="f-ring"></div><div class="tbl" id="f-tabel"></div>';
  const load = jaga(async () => {
    const d = await api('/api/penjualan?' + qs(nilaiFilter('f')));
    const r = d.ringkas;
    $('#f-ring').innerHTML = `<div><b>${num(r.jumlah)}</b><span>entri</span></div>
      <div><b>${num(r.botol_jual)}</b><span>botol terjual</span></div>
      <div><b>${num(r.botol_sample)}</b><span>botol sample</span></div>
      <div><b>${rp(r.total)}</b><span>total rupiah</span></div>
      <div><b>${num(r.toko)}</b><span>toko</span></div><div><b>${num(r.spg)}</b><span>SPG</span></div>`;
    $('#f-tabel').innerHTML = d.rows.length ? `<table><thead><tr><th>Tanggal</th><th>Jam</th><th>SPG</th><th>Toko</th>
      <th>Produk</th><th>Jenis</th><th>Qty</th><th>Satuan</th><th>Botol</th><th>Harga</th><th>Total</th><th>Catatan</th><th></th></tr></thead>
      <tbody>${d.rows.map(e => `<tr><td>${tgl(e.tanggal)}</td><td>${jam(e.waktu)}</td><td>${esc(e.spg)}</td><td>${esc(e.toko)}</td>
        <td>${esc(e.produk)}</td><td>${e.jenis === 'SAMPLE' ? '<span class="badge b-sample">SAMPLE</span>' : '<span class="badge b-jual">JUAL</span>'}</td>
        <td class="n">${num(e.qty)}</td><td>${e.satuan}</td><td class="n">${num(e.botol_setara)}</td><td class="n">${num(e.harga)}</td>
        <td class="n">${num(e.total)}</td><td>${esc(e.catatan)}</td>
        <td><button class="btn sm danger" data-hapus="${e.uid}">Hapus</button></td></tr>`).join('')}</tbody></table>`
      : '<div class="empty">Tidak ada data pada filter ini.</div>';
  });
  $('#f-go').onclick = load;
  $('#f-q').addEventListener('keydown', e => e.key === 'Enter' && load());
  ['spg', 'toko', 'produk', 'jenis'].forEach(k => $(`#f-${k}`).onchange = load);
  $('#f-xls').onclick = () => { unduh('/export/penjualan?' + qs(nilaiFilter('f'))); toast('Menyiapkan file…'); };
  pasangCepat(el, 'f', load);
  el.addEventListener('click', jaga(async e => {
    const b = e.target.closest('[data-hapus]'); if (!b) return;
    if (!await konfirmasi('Hapus data', 'Hapus entri penjualan ini? Tindakan tercatat di log.', 'Hapus')) return;
    await api('/api/penjualan/' + b.dataset.hapus, {method: 'DELETE'}); toast('Dihapus'); load();
  }));
  load();
}

// ---------- laporan (rekap)
async function mLaporan(el) {
  const opsi = await opsiFilter();
  el.innerHTML = filterBar('l', opsi, '<button class="btn" id="l-xls">Export Excel</button>') +
    '<div class="ring" id="l-ring"></div><div class="grid2" id="l-grup"></div>';
  const judul = {spg: 'Per SPG', toko: 'Per Toko', produk: 'Per Produk', tanggal: 'Per Tanggal', wilayah: 'Per Wilayah'};
  const load = jaga(async () => {
    const d = await api('/api/laporan?' + qs({...nilaiFilter('l'), limit: 1}));
    const r = d.ringkas;
    $('#l-ring').innerHTML = `<div><b>${num(r.jumlah)}</b><span>entri</span></div>
      <div><b>${num(r.botol_jual)}</b><span>botol terjual</span></div>
      <div><b>${num(r.botol_sample)}</b><span>botol sample</span></div>
      <div><b>${rp(r.total)}</b><span>total rupiah</span></div>`;
    $('#l-grup').innerHTML = Object.entries(d.grup).map(([k, arr]) => `<div class="card"><h2>${judul[k]}</h2>
      <div class="tbl" style="max-height:340px"><table><thead><tr><th>${judul[k].replace('Per ', '')}</th><th>Entri</th><th>Botol</th><th>Sample</th><th>Total</th></tr></thead>
      <tbody>${arr.map(x => `<tr><td>${esc(k === 'tanggal' ? tgl(x.label) : x.label)}</td><td class="n">${x.entri}</td>
        <td class="n">${num(x.botol_jual)}</td><td class="n">${num(x.botol_sample)}</td><td class="n">${num(x.total)}</td></tr>`).join('')
        || '<tr><td colspan="5" class="empty">Tidak ada data</td></tr>'}</tbody></table></div></div>`).join('');
  });
  $('#l-go').onclick = load;
  ['spg', 'toko', 'produk', 'jenis'].forEach(k => $(`#l-${k}`).onchange = load);
  $('#l-xls').onclick = () => unduh('/export/penjualan?' + qs(nilaiFilter('l')));
  pasangCepat(el, 'l', load);
  load();
}

// ---------- akun SPG
async function mSpg(el) {
  const [users, toko] = await Promise.all([api('/api/pengguna'), api('/api/toko')]);
  el.innerHTML = `<div class="card"><div class="row"><h2 style="margin:0;flex:1">Akun pengguna</h2>
      <button class="btn primary" id="u-baru">+ Akun baru</button></div></div>
    <div class="tbl"><table><thead><tr><th>Nama</th><th>Username</th><th>Role</th><th>Toko</th><th>Entri</th>
      <th>Status</th><th>Terakhir masuk</th><th>Aksi</th></tr></thead><tbody>
      ${users.map(u => `<tr><td><b>${esc(u.nama)}</b></td><td><code>${esc(u.username)}</code></td>
        <td>${u.role === 'master' ? '<span class="badge b-jual">MASTER</span>' : 'SPG'}</td>
        <td>${u.role === 'spg' ? `<button class="btn sm" data-toko="${u.uid}">${u.jumlah_toko} toko</button>` : '-'}</td>
        <td class="n">${num(u.jumlah_entri)}</td>
        <td>${!u.aktif ? '<span class="badge b-err">nonaktif</span>' : u.kunci_sampai > new Date().toISOString() ? '<span class="badge b-err">terkunci</span>' : '<span class="badge b-ok">aktif</span>'}
            ${u.totp_aktif ? ' <span class="badge b-ok">2FA</span>' : ''}${u.harus_ganti ? ' <span class="badge b-sample">password baru</span>' : ''}</td>
        <td>${u.terakhir_login ? tgl(u.terakhir_login) + ' ' + jam(u.terakhir_login) : '-'}</td>
        <td><button class="btn sm" data-ubah="${u.uid}">Ubah</button>
            <button class="btn sm" data-reset="${u.uid}">Reset password</button>
            ${u.kunci_sampai > new Date().toISOString() ? `<button class="btn sm" data-buka="${u.uid}">Buka kunci</button>` : ''}
            <button class="btn sm danger" data-hapus="${u.uid}">Hapus</button></td></tr>`).join('')}
    </tbody></table></div>`;
  const form = (u = null) => {
    const w = document.createElement('div');
    w.innerHTML = `<div class="form">
      <div class="full"><label class="f">Nama lengkap *</label><input name="nama" value="${esc(u?.nama || '')}"></div>
      <div><label class="f">Username</label><input name="username" value="${esc(u?.username || '')}" placeholder="otomatis dari nama"></div>
      <div><label class="f">Role</label><select name="role"><option value="spg" ${u?.role !== 'master' ? 'selected' : ''}>SPG</option>
        <option value="master" ${u?.role === 'master' ? 'selected' : ''}>Master</option></select></div>
      <div><label class="f">Telepon</label><input name="telepon" value="${esc(u?.telepon || '')}"></div>
      <div><label class="f">Status</label><select name="aktif"><option value="1" ${u?.aktif !== 0 ? 'selected' : ''}>Aktif</option>
        <option value="0" ${u?.aktif === 0 ? 'selected' : ''}>Nonaktif</option></select></div>
      <div class="full"><label class="f">Catatan</label><input name="catatan" value="${esc(u?.catatan || '')}"></div>
      <div class="full"><label class="f">Toko yang ditugaskan</label>
        <div class="checklist">${toko.map(t => `<label><input type="checkbox" value="${t.uid}"
          ${u?.toko?.some(x => x.uid === t.uid) ? 'checked' : ''}>${esc(t.nama)} <span class="muted">${esc(t.wilayah || '')}</span></label>`).join('')}</div></div>
    </div>`;
    buka(u ? 'Ubah Akun' : 'Akun Baru', w, [{teks: 'Batal'}, {teks: 'Simpan', cls: 'primary', klik: async () => {
      const body = Object.fromEntries($$('[name]', w).map(i => [i.name, i.value]));
      body.aktif = body.aktif === '1';
      body.toko_uids = $$('.checklist input:checked', w).map(i => i.value);
      const r = u ? await api('/api/pengguna/' + u.uid, {method: 'PUT', body}) : await api('/api/pengguna', {method: 'POST', body});
      if (r.password) {
        buka('Akun dibuat', `<div class="note ok">Catat dan berikan ke SPG. Password hanya ditampilkan sekali.</div>
          <p>Username: <code>${esc(r.username)}</code></p><p>Password: <code style="font-size:18px">${esc(r.password)}</code></p>
          <p class="muted">SPG wajib mengganti password saat pertama kali masuk.</p>`, [{teks: 'Sudah dicatat', cls: 'primary'}]);
      } else toast('Tersimpan');
      mSpg(el);
      return !r.password;
    }}]);
  };
  $('#u-baru').onclick = () => form();
  el.addEventListener('click', jaga(async e => {
    const b = e.target.closest('button'); if (!b) return;
    const u = users.find(x => x.uid === (b.dataset.ubah || b.dataset.reset || b.dataset.hapus || b.dataset.buka || b.dataset.toko));
    if (b.dataset.ubah || b.dataset.toko) form(u);
    if (b.dataset.buka) { await api(`/api/pengguna/${u.uid}/buka`, {method: 'POST', body: {}}); toast('Kunci dibuka'); mSpg(el); }
    if (b.dataset.reset) {
      if (!await konfirmasi('Reset password', `Buat password baru untuk <b>${esc(u.nama)}</b>? Semua sesi akan keluar.`, 'Reset', false)) return;
      const r = await api(`/api/pengguna/${u.uid}/reset`, {method: 'POST', body: {}});
      buka('Password baru', `<div class="note ok">Berikan ke ${esc(u.nama)}. Hanya ditampilkan sekali.</div>
        <p>Username: <code>${esc(r.username)}</code></p><p>Password: <code style="font-size:18px">${esc(r.password)}</code></p>`,
        [{teks: 'Sudah dicatat', cls: 'primary'}]);
      mSpg(el);
    }
    if (b.dataset.hapus) {
      if (!await konfirmasi('Hapus akun', `Hapus akun <b>${esc(u.nama)}</b>? Jika sudah punya data, akun hanya dinonaktifkan.`, 'Hapus')) return;
      const r = await api('/api/pengguna/' + u.uid, {method: 'DELETE'});
      toast(r.dinonaktifkan ? 'Akun dinonaktifkan (punya data)' : 'Akun dihapus'); mSpg(el);
    }
  }));
}

// ---------- toko & produk
async function mToko(el) {
  const [toko, users] = await Promise.all([api('/api/toko?aktif=0'), api('/api/pengguna')]);
  const spg = users.filter(u => u.role === 'spg');
  el.innerHTML = `<div class="card"><div class="row"><h2 style="margin:0;flex:1">Daftar toko</h2>
      <input id="t-cari" placeholder="Cari toko…" style="max-width:260px"><button class="btn primary" id="t-baru">+ Toko baru</button></div></div>
    <div class="tbl" id="t-tabel"></div>`;
  const gambar = () => {
    const q = $('#t-cari').value.trim().toUpperCase();
    const list = toko.filter(t => !q || (t.nama + ' ' + (t.wilayah || '') + ' ' + (t.alamat || '')).toUpperCase().includes(q));
    $('#t-tabel').innerHTML = `<table><thead><tr><th>Nama</th><th>Wilayah</th><th>Alamat</th><th>SPG</th><th>Entri</th><th>Status</th><th></th></tr></thead>
      <tbody>${list.map(t => `<tr><td><b>${esc(t.nama)}</b></td><td>${esc(t.wilayah || '')}</td><td>${esc(t.alamat || '')}</td>
        <td>${esc(t.spg || '-')}</td><td class="n">${num(t.jumlah_entri)}</td>
        <td>${t.aktif ? '<span class="badge b-ok">aktif</span>' : '<span class="badge b-err">nonaktif</span>'}</td>
        <td><button class="btn sm" data-ubah="${t.uid}">Ubah</button> <button class="btn sm danger" data-hapus="${t.uid}">Hapus</button></td></tr>`).join('')
        || '<tr><td colspan="7" class="empty">Belum ada toko.</td></tr>'}</tbody></table>`;
  };
  const form = (t = null) => {
    const w = document.createElement('div');
    w.innerHTML = `<div class="form">
      <div class="full"><label class="f">Nama toko *</label><input name="nama" value="${esc(t?.nama || '')}"></div>
      <div><label class="f">Wilayah</label><input name="wilayah" value="${esc(t?.wilayah || '')}"></div>
      <div><label class="f">Status</label><select name="aktif"><option value="1" ${t?.aktif !== 0 ? 'selected' : ''}>Aktif</option>
        <option value="0" ${t?.aktif === 0 ? 'selected' : ''}>Nonaktif</option></select></div>
      <div class="full"><label class="f">Alamat</label><input name="alamat" value="${esc(t?.alamat || '')}"></div>
      <div class="full"><label class="f">Catatan</label><input name="catatan" value="${esc(t?.catatan || '')}"></div>
      <div class="full"><label class="f">SPG yang bertugas</label><div class="checklist">${spg.map(u => `<label>
        <input type="checkbox" value="${u.uid}" ${u.toko.some(x => x.uid === t?.uid) ? 'checked' : ''}>${esc(u.nama)}</label>`).join('')
        || '<span class="muted">Belum ada akun SPG.</span>'}</div></div></div>`;
    buka(t ? 'Ubah Toko' : 'Toko Baru', w, [{teks: 'Batal'}, {teks: 'Simpan', cls: 'primary', klik: async () => {
      const body = Object.fromEntries($$('[name]', w).map(i => [i.name, i.value]));
      body.aktif = body.aktif === '1';
      body.spg_uids = $$('.checklist input:checked', w).map(i => i.value);
      t ? await api('/api/toko/' + t.uid, {method: 'PUT', body}) : await api('/api/toko', {method: 'POST', body});
      toast('Tersimpan'); mToko(el);
    }}]);
  };
  $('#t-baru').onclick = () => form();
  $('#t-cari').oninput = gambar;
  el.addEventListener('click', jaga(async e => {
    const b = e.target.closest('button'); if (!b) return;
    const t = toko.find(x => x.uid === (b.dataset.ubah || b.dataset.hapus));
    if (b.dataset.ubah) form(t);
    if (b.dataset.hapus) {
      if (!await konfirmasi('Hapus toko', `Hapus <b>${esc(t.nama)}</b>? Jika sudah ada data penjualan, toko hanya dinonaktifkan.`, 'Hapus')) return;
      const r = await api('/api/toko/' + t.uid, {method: 'DELETE'});
      toast(r.dinonaktifkan ? 'Toko dinonaktifkan' : 'Toko dihapus'); mToko(el);
    }
  }));
  gambar();
}

async function mProduk(el) {
  const produk = await api('/api/produk?semua=1');
  el.innerHTML = `<div class="card"><div class="row"><h2 style="margin:0;flex:1">Daftar produk & harga</h2>
      <button class="btn primary" id="p-baru">+ Produk baru</button></div>
      <p class="muted" style="margin:8px 0 0;font-size:14px">Harga di sini muncul otomatis di HP SPG dan masih bisa diubah saat input bila ada perubahan harga.</p></div>
    <div class="tbl"><table><thead><tr><th>Urut</th><th>Kode</th><th>Nama</th><th>Harga / botol</th><th>Harga / dus</th><th>Isi per dus</th><th>Entri</th><th>Status</th><th></th></tr></thead>
      <tbody>${produk.map(p => `<tr><td class="n">${p.urut}</td><td><code>${esc(p.kode)}</code></td><td><b>${esc(p.nama)}</b></td>
        <td class="n">${num(p.harga_botol)}</td><td class="n">${num(p.harga_dus)}</td><td class="n">${p.isi_dus}</td>
        <td class="n">${num(p.jumlah_entri)}</td>
        <td>${p.aktif ? '<span class="badge b-ok">aktif</span>' : '<span class="badge b-err">nonaktif</span>'}</td>
        <td><button class="btn sm" data-ubah="${p.uid}">Ubah</button> <button class="btn sm danger" data-hapus="${p.uid}">Hapus</button></td></tr>`).join('')
        || '<tr><td colspan="9" class="empty">Belum ada produk. Klik "+ Produk baru".</td></tr>'}</tbody></table></div>`;
  const form = (p = null) => {
    const w = document.createElement('div');
    w.innerHTML = `<div class="form">
      <div><label class="f">Kode *</label><input name="kode" value="${esc(p?.kode || '')}"></div>
      <div class="full"><label class="f">Nama produk *</label><input name="nama" value="${esc(p?.nama || '')}"></div>
      <div><label class="f">Harga per botol</label><input name="harga_botol" inputmode="numeric" value="${p?.harga_botol || ''}"></div>
      <div><label class="f">Harga per dus</label><input name="harga_dus" inputmode="numeric" value="${p?.harga_dus || ''}"></div>
      <div><label class="f">Isi per dus</label><input name="isi_dus" inputmode="numeric" value="${p?.isi_dus ?? 12}"></div>
      <div><label class="f">Urutan tampil</label><input name="urut" inputmode="numeric" value="${p?.urut ?? 0}"></div>
      <div><label class="f">Status</label><select name="aktif"><option value="1" ${p?.aktif !== 0 ? 'selected' : ''}>Aktif</option>
        <option value="0" ${p?.aktif === 0 ? 'selected' : ''}>Nonaktif</option></select></div></div>
      <p class="muted" style="font-size:13px">Kosongkan harga dus bila ingin dihitung otomatis dari harga botol × isi per dus.</p>`;
    buka(p ? 'Ubah Produk' : 'Produk Baru', w, [{teks: 'Batal'}, {teks: 'Simpan', cls: 'primary', klik: async () => {
      const body = Object.fromEntries($$('[name]', w).map(i => [i.name, i.value]));
      body.aktif = body.aktif === '1';
      p ? await api('/api/produk/' + p.uid, {method: 'PUT', body}) : await api('/api/produk', {method: 'POST', body});
      toast('Tersimpan'); mProduk(el);
    }}]);
  };
  $('#p-baru').onclick = () => form();
  el.addEventListener('click', jaga(async e => {
    const b = e.target.closest('button'); if (!b) return;
    const p = produk.find(x => x.uid === (b.dataset.ubah || b.dataset.hapus));
    if (b.dataset.ubah) form(p);
    if (b.dataset.hapus) {
      if (!await konfirmasi('Hapus produk', `Hapus <b>${esc(p.nama)}</b>? Jika sudah dipakai, produk hanya dinonaktifkan.`, 'Hapus')) return;
      const r = await api('/api/produk/' + p.uid, {method: 'DELETE'});
      toast(r.dinonaktifkan ? 'Produk dinonaktifkan' : 'Produk dihapus'); mProduk(el);
    }
  }));
}

// ---------- log keamanan
async function mLog(el) {
  el.innerHTML = `<div class="card"><div class="form">
      <div><label class="f">Dari</label><input type="date" id="g-dari"></div>
      <div><label class="f">Sampai</label><input type="date" id="g-sampai"></div>
      <div><label class="f">Aktivitas</label><select id="g-aksi"><option value="">Semua</option></select></div>
      <div><label class="f">Cari</label><input id="g-q" placeholder="username / IP / detail"></div>
      <div class="full"><button class="btn primary" id="g-go">Tampilkan</button></div></div></div>
    <div class="tbl" id="g-tabel"></div>`;
  let pertama = true;
  const load = jaga(async () => {
    const d = await api('/api/log?' + qs({dari: $('#g-dari').value, sampai: $('#g-sampai').value, aksi: $('#g-aksi').value, q: $('#g-q').value}));
    if (pertama) { $('#g-aksi').innerHTML += d.aksi.map(a => `<option value="${a}">${a}</option>`).join(''); pertama = false; }
    $('#g-tabel').innerHTML = d.rows.length ? `<table><thead><tr><th>Waktu</th><th>Pengguna</th><th>Role</th><th>Aktivitas</th><th>Detail</th><th>IP</th><th>Perangkat</th></tr></thead>
      <tbody>${d.rows.map(l => `<tr><td>${tgl(l.waktu)} ${jam(l.waktu)}</td><td>${esc(l.username)}</td><td>${esc(l.role)}</td>
        <td><span class="badge ${/GAGAL|HAPUS|KUNCI/.test(l.aksi) ? 'b-err' : 'b-jual'}">${esc(l.aksi)}</span></td>
        <td>${esc(l.detail)}</td><td>${esc(l.ip)}</td><td>${esc((l.perangkat || '').slice(0, 45))}</td></tr>`).join('')}</tbody></table>`
      : '<div class="empty">Tidak ada log.</div>';
  });
  $('#g-go').onclick = load; $('#g-q').addEventListener('keydown', e => e.key === 'Enter' && load());
  $('#g-aksi').onchange = load;
  load();
}

// ============================================================ start
(async () => {
  try {
    const me = await api('/api/me');
    if (me.login) { ME = me; await mulai(); } else tampilLogin();
  } catch { tampilLogin(); }
})();

</script>
</body>
</html>
'''


def main():
    init_db()
    backup_harian()
    if HOST not in ("127.0.0.1", "localhost") and not TRUST_PROXY and not DEV_MODE:
        print("PERINGATAN: aplikasi dibuka ke jaringan tanpa reverse proxy HTTPS.")
        print("Untuk online, jalankan di belakang Caddy/nginx (lihat PANDUAN_DEPLOY.txt) lalu set SPG_TRUST_PROXY=1.")
    server = ThreadingHTTPServer((HOST, PORT), Handler)
    server.daemon_threads = True
    print("=" * 64)
    print(f"  {APP_NAME} v{APP_VERSION}")
    print(f"  Alamat lokal : http://{HOST}:{PORT}")
    print(f"  Data         : {DB_PATH}")
    print(f"  Mode         : {'DEV (HTTPS tidak dipaksa)' if DEV_MODE else 'PRODUKSI'}"
          f"{' | di belakang proxy' if TRUST_PROXY else ''}")
    print("=" * 64)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nBerhenti.")
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
